import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import app


def test_authored_html_exports_real_pptx_with_editable_text():
    page = '<style>body{margin:0;background:#fff}h1{position:absolute;left:80px;top:50px;font:40px Arial;color:#123456}svg{position:absolute;left:300px;top:200px}</style><h1 data-editable="heading">Direct lesson</h1><svg width="400" height="200"><path d="M0 180 Q200 -100 400 180" fill="none" stroke="blue" stroke-width="8"/></svg>'
    response = TestClient(app).post('/render/pptx', json={
        'title': 'Direct', 'theme': 'classroom',
        'slides': [{'title': 'Direct lesson', 'html': page, 'notes': 'Teacher notes'}],
    })
    assert response.status_code == 200, response.text[:500] if response.status_code != 200 else ''
    ppt = Presentation(io.BytesIO(response.content))
    texts = [s for s in ppt.slides[0].shapes if s.has_text_frame]
    assert any(s.text == 'Direct lesson' for s in texts)
    heading = next(s for s in texts if s.text == 'Direct lesson')
    assert abs(heading.left / ppt.slide_width - 80 / 1280) < .01
    assert any(s.shape_type == 13 for s in ppt.slides[0].shapes)
    assert 'Teacher notes' in ppt.slides[0].notes_slide.notes_text_frame.text
    assert zipfile.is_zipfile(io.BytesIO(response.content))


@pytest.mark.parametrize('source', [
    '<script>fetch("http://127.0.0.1:1/secret")</script><p data-editable="x">Blocked</p>',
    '<img src="file:///C:/Windows/win.ini"><p data-editable="x">Blocked</p>',
    '<script>while(true){}</script>',
])
def test_authored_page_fails_closed_on_external_resource_or_hang(source):
    response = TestClient(app).post('/render/pptx', json={
        'title': 'Direct', 'slides': [{'title': 'Unsafe', 'html': source}],
    })
    assert response.status_code == 422



def test_authored_picture_and_rich_text_are_not_dropped():
    import base64
    from PIL import Image
    image=io.BytesIO();Image.new('RGB',(40,40),(255,0,0)).save(image,format='PNG')
    source='<style>body{margin:0}p{position:absolute;left:80px;top:100px;font:32px Arial}b{color:rgb(255,0,0)}img{position:absolute;left:400px;top:100px;width:100px;height:100px}</style><p data-editable="x">Hello <b>world</b></p><img data-teaching-picture>'
    response=TestClient(app).post('/render/pptx',json={'title':'Direct','slides':[{'title':'Rich','html':source,'picture':{'data_url':'data:image/png;base64,'+base64.b64encode(image.getvalue()).decode(),'caption':'Red square','source':'Teacher'}}]})
    assert response.status_code==200,response.text[:500] if response.status_code!=200 else ''
    ppt=Presentation(io.BytesIO(response.content))
    runs=[r for s in ppt.slides[0].shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs]
    assert any(r.text=='world' and r.font.bold and str(r.font.color.rgb)=='FF0000' for r in runs)
    artwork=next(s for s in ppt.slides[0].shapes if s.name=='authored-artwork')
    assert Image.open(io.BytesIO(artwork.image.blob)).convert('RGB').getpixel((450,150))==(255,0,0)


def test_authored_page_rejects_unmarked_visible_text():
    response=TestClient(app).post('/render/pptx',json={'title':'Direct','slides':[{'title':'Missing','html':'<p data-editable="x">Editable</p><p>Unmarked words</p>'}]})
    assert response.status_code==422
    assert response.json()['detail']['code']=='authored_text_unmarked'
    assert response.json()['detail']['field']=='html'


def test_authored_page_rejects_overlapping_labels():
    source='<style>p{position:absolute;left:100px;top:100px;font:32px Arial}</style><p data-editable="a">Overlapping heading</p><p data-editable="b">Misplaced label</p>'
    result=TestClient(app).post('/render/pptx',json={'title':'Direct','slides':[{'title':'Overlap','html':source}]})
    assert result.status_code==422
    assert result.json()['detail']['code']=='authored_text_overlap'
    assert result.json()['detail']['page']==1


@pytest.mark.parametrize(('source', 'code'), [
    ('<p data-editable="x" style="position:absolute;left:1250px;width:200px">Outside</p>', 'authored_text_outside'),
    ('<p data-editable="x" style="width:30px;height:20px;font:32px Arial;overflow:hidden">Too much text</p>', 'authored_text_overflow'),
    ('<p data-editable="x" style="width:200px;transform:translateX(10px)">Moved</p>', 'authored_text_transform'),
    ('<p data-editable="x">First</p><p data-editable="x">Duplicate</p>', 'authored_text_ids'),
])
def test_authored_page_returns_safe_specific_repair_category(source, code):
    result = TestClient(app).post('/render/pptx', json={'title':'Direct','slides':[{'title':'Repair','html':source}]})
    assert result.status_code == 422
    assert result.json()['detail']['code'] == code
    assert 'Outside' not in result.text and 'Too much text' not in result.text
