import base64
import io
import zipfile

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from app import app


def test_parameter_animation_is_embedded_as_multiframe_gif():
    graph = dict(title='参数a与开口方向', x_min=-4, x_max=4, y_min=-5, y_max=5,
                 series=[dict(kind='quadratic', a=1, b=0, c=0)],
                 animation=dict(series=0, param='a', to=-1, duration_ms=2000))
    result = TestClient(app).post('/render/pptx', json=dict(title='函数动图', slides=[dict(
        title='改变a观察开口', graph=graph, bullets=['a为零时不是二次函数。'])]))
    assert result.status_code == 200, result.text
    ppt = Presentation(io.BytesIO(result.content))
    assert len(ppt.slides) == 1
    with zipfile.ZipFile(io.BytesIO(result.content)) as z:
        names = [n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.gif')]
        assert len(names) == 1, 'parameter animation must survive the real PPTX export'
        with Image.open(io.BytesIO(z.read(names[0]))) as im:
            assert im.n_frames >= 20
            assert im.info['loop'] == 0
            im.seek(0); first = im.convert('RGB').tobytes()
            im.seek(im.n_frames-1)
            assert im.convert('RGB').tobytes() != first


def test_inverse_animation_cannot_cross_zero_coefficient():
    graph = dict(title='反比例', x_min=-4, x_max=4, y_min=-5, y_max=5,
                 series=[dict(kind='inverse', k=1)], animation=dict(series=0, param='k', to=-1, duration_ms=2000))
    r = TestClient(app).post('/render/pptx', json=dict(title='函数动图', slides=[dict(title='定义域', graph=graph)]))
    assert r.status_code == 422


def test_resolved_image_references_preserve_gif_and_large_photo_bytes():
    import random
    frames = [Image.new('RGB', (64,64), c) for c in ('red','blue')]
    output = io.BytesIO();frames[0].save(output,format='GIF',save_all=True,append_images=frames[1:],duration=200,loop=0)
    gif = output.getvalue()
    photo = io.BytesIO();Image.frombytes('RGB',(400,400),random.Random(7).randbytes(400*400*3)).save(photo,format='PNG')
    assert len(photo.getvalue()) > 128000
    pictures = [('gif',gif),('png',photo.getvalue())]
    result = TestClient(app).post('/render/pptx', json=dict(title='引用图片', slides=[dict(
        title='图片与出处', picture=dict(asset_id='a'*32,data_url='data:image/'+kind+';base64,'+base64.b64encode(raw).decode(),caption='教师配图',source='教师本人提供')) for kind,raw in pictures]))
    assert result.status_code == 200, result.text
    with zipfile.ZipFile(io.BytesIO(result.content)) as z:
        blobs=[z.read(n) for n in z.namelist() if n.startswith('ppt/media/')]
        assert all(raw in blobs for _,raw in pictures)
    ppt=Presentation(io.BytesIO(result.content))
    assert all('教师本人提供' in s.notes_slide.notes_text_frame.text for s in ppt.slides)


def test_function_graph_and_comparison_table_both_survive_export():
    graph=dict(title='读图并比较',x_min=-3,x_max=3,y_min=-3,y_max=3,series=[dict(kind='quadratic',a=1,b=0,c=0)])
    table=dict(headers=['参数','开口'],rows=[['a > 0','向上'],['a < 0','向下']])
    result=TestClient(app).post('/render/pptx',json=dict(title='函数图与表',slides=[dict(
        title='参数与图像',graph=graph,table=table,bullets=['结合图像解释表中的比较。'])]))
    assert result.status_code==200,result.text
    ppt=Presentation(io.BytesIO(result.content))
    assert any(s.has_table for s in ppt.slides[0].shapes)
    assert any(s.name.startswith('function-') for s in ppt.slides[0].shapes)
    from test_render import assert_inside_slide
    assert_inside_slide(ppt)
