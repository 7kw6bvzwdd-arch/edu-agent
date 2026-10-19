import base64
import io

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from app import app


import pytest


@pytest.mark.parametrize('provider,label', [('wikimedia_commons', 'Wikimedia Commons'), ('web_page', '网页配图')])
@pytest.mark.parametrize('placement', ['picture', 'canvas'])
def test_image_exports_full_source_snapshot_in_real_pptx(placement, provider, label):
    image = io.BytesIO()
    Image.new('RGB', (320, 240), 'tan').save(image, format='PNG')
    provenance = {
        'provider': provider, 'title': 'File:Great Wall.jpg',
        'page_url': 'https://commons.wikimedia.org/wiki/File:Great_Wall.jpg',
        'original_url': 'https://upload.wikimedia.org/wikipedia/commons/a/ab/Great_Wall.jpg',
        'download_url': 'https://upload.wikimedia.org/wikipedia/commons/a/ab/Great_Wall.jpg',
        'author': 'Test Photographer', 'license': 'CC BY-SA 4.0',
        'license_url': 'https://creativecommons.org/licenses/by-sa/4.0/',
        'description': 'Great Wall at Badaling', 'date': '2020-01-01',
        'retrieved_at': '2026-10-19T06:00:00Z', 'sha256': 'a' * 64,
    }
    payload = {
        'title': '观察长城', 'theme': 'ocean', 'slides': [{
            'title': '观察长城防御设施', 'bullets': ['观察城墙与敌台的连接。'],
            'picture': {'asset_id': 'a' * 32, 'caption': '长城的防御设施',
                        'source': 'Wikimedia Commons · CC BY-SA 4.0',
                        'data_url': 'data:image/png;base64,' + base64.b64encode(image.getvalue()).decode(),
                        'provenance': provenance},
        }],
    }
    if placement == 'canvas':
        page = payload['slides'][0]
        picture = page.pop('picture')
        page['canvas_objects'] = [dict(picture, id='object-source', kind='image', x=.2, y=.3, width=.4, height=.4)]
    response = TestClient(app).post('/render/pptx', json=payload)
    assert response.status_code == 200, response.text
    deck = Presentation(io.BytesIO(response.content))
    assert any(shape.shape_type == 13 for shape in deck.slides[0].shapes)
    notes = deck.slides[0].notes_slide.notes_text_frame.text
    for expected in ['图片来源（' + label + '）', 'Test Photographer',
                     provenance['page_url'], provenance['license_url'],
                     provenance['sha256'], provenance['retrieved_at']]:
        assert expected in notes
    assert '教师填写' not in notes
