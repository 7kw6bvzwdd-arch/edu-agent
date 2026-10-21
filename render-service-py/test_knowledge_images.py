import base64
import io

import fitz
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app import app


def image_bytes():
    output = io.BytesIO()
    Image.new('RGB', (320, 240), '#246891').save(output, 'PNG')
    return output.getvalue()


def test_uploaded_pdf_and_grouped_pptx_extract_reusable_images_with_page_sources():
    raw = image_bytes()
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((40, 40), 'Water cycle diagram')
    page.insert_image(fitz.Rect(40, 60, 360, 300), stream=raw)
    pdf_data = pdf.tobytes(); pdf.close()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(.1), Inches(4), Inches(1)).text = 'Water cycle diagram'
    group = slide.shapes.add_group_shape()
    group.shapes.add_picture(io.BytesIO(raw), Inches(1), Inches(1))
    output = io.BytesIO(); deck.save(output)
    client = TestClient(app)
    for name, data in [('cycle.pdf', pdf_data), ('cycle.pptx', output.getvalue())]:
        response = client.post('/extract', files={'file': (name, data)})
        assert response.status_code == 200, response.text
        images = response.json().get('images', [])
        assert len(images) == 1, 'embedded figure must be extracted rather than a page screenshot'
        figure = images[0]
        assert figure['page'] == 1
        assert 'Water cycle diagram' in figure['caption']
        with Image.open(io.BytesIO(base64.b64decode(figure['data_url'].split(',', 1)[1]))) as im:
            assert im.size == (320, 240)
            assert im.convert('RGB').getpixel((100, 100)) == (36, 104, 145)


def test_image_only_slide_is_indexable_and_extraction_is_bounded():
    deck = Presentation()
    for i in range(28):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(image_bytes()), Inches(1), Inches(1))
    output = io.BytesIO(); deck.save(output)
    response = TestClient(app).post('/extract', files={'file': ('images.pptx', output.getvalue())})
    assert response.status_code == 200, response.text
    assert len(response.json()['images']) == 24
    assert response.json()['image_warnings']
    assert len(response.json()['pages']) == 28


def test_pdf_invalid_image_mask_preserves_text_and_other_images():
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((40, 40), 'Readable textbook content')
    broken = page.insert_image(fitz.Rect(40, 60, 360, 300), stream=image_bytes())
    mask = fitz.Pixmap(fitz.csGRAY, fitz.IRect(0, 0, 80, 80), False)
    mask.clear_with(255)
    mask_xref = page.insert_image(fitz.Rect(400, 60, 480, 140), pixmap=mask)
    pdf.xref_set_key(broken, 'SMask', f'{mask_xref} 0 R')
    output = io.BytesIO()
    Image.new('RGB', (320, 240), '#FF3300').save(output, 'PNG')
    page.insert_image(fitz.Rect(40, 340, 360, 580), stream=output.getvalue())
    data = pdf.tobytes(); pdf.close()
    response = TestClient(app).post('/extract', files={'file': ('broken-mask.pdf', data)})
    assert response.status_code == 200, response.text
    assert 'Readable textbook content' in response.json()['pages'][0]['text']
    assert response.json()['image_warnings'], 'unsupported image must remain visible as a warning'
    assert response.json()['images'], 'other valid images must remain available'
