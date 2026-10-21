import io
import base64
import json
import zipfile

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from app import app


def test_moving_shape_label_follows_in_rendered_frames_and_pptx():
    diagram=dict(kind='scene',title='物体与标注同步',scene=dict(duration_ms=2000,
        objects=[dict(id='boat',type='polygon',points=[[.2,.35],[.4,.35],[.36,.42],[.24,.42]],fill='#1767bf'),
                 dict(id='label',type='text',x=.3,y=.53,text='小船',attach_to='boat',color='#c00000')],
        actions=[dict(object='boat',op='move',x=.7,y=.385,start_ms=0,duration_ms=2000)]))
    client=TestClient(app)
    response=client.post('/render/scientific',json=dict(diagram=diagram,frames=True))
    assert response.status_code==200,response.text
    frames=response.json()['frames']
    # Inspect the actual red label pixels; a stationary or absent label fails.
    centers=[]
    for frame in (frames[0],frames[-1]):
        im=Image.open(io.BytesIO(base64.b64decode(frame.split(',')[1]))).convert('RGB')
        pixels=[x for y in range(310,370) for x in range(30,970) if (lambda p:p[0]>140 and p[1]<80 and p[2]<80)(im.getpixel((x,y)))]
        assert pixels
        centers.append(sum(pixels)/len(pixels))
    assert 370<centers[1]-centers[0]<382
    ppt=client.post('/render/pptx',json=dict(title='同步标注',slides=[dict(title='移动图示',scientific=diagram)]))
    assert ppt.status_code==200,ppt.text
    with zipfile.ZipFile(io.BytesIO(ppt.content)) as z:
        gif=next(n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.gif'))
        with Image.open(io.BytesIO(z.read(gif))) as im:assert im.n_frames>1


def test_route_cannot_teleport_from_its_declared_start():
    diagram=dict(kind='scene',title='路线方向',scene=dict(duration_ms=2000,
        objects=[dict(id='route',type='path',points=[[.2,.5],[.8,.5]]),
                 dict(id='ball',type='ellipse',x=.8,y=.5,width=.05,height=.05)],
        actions=[dict(object='ball',op='follow_path',path='route',start_ms=0,duration_ms=2000)]))
    client=TestClient(app)
    rejected=client.post('/render/scientific',json=dict(diagram=diagram,frames=True))
    assert rejected.status_code==422
    diagram['scene']['objects'][0]['points'].reverse()
    accepted=client.post('/render/scientific',json=dict(diagram=diagram,frames=True))
    assert accepted.status_code==200,accepted.text
    assert accepted.json()['frames'][0]!=accepted.json()['frames'][-1]


def test_highlight_background_never_covers_teaching_label():
    diagram=dict(kind='scene',title='标注与高亮',scene=dict(duration_ms=1000,
        objects=[dict(id='word',type='text',x=.5,y=.5,text='本体',color='#c00000'),
                 dict(id='highlight',type='rect',x=.5,y=.5,width=.3,height=.15,fill='#ffff99')]))
    response=TestClient(app).post('/render/scientific',json=dict(diagram=diagram))
    assert response.status_code==200,response.text
    im=Image.open(io.BytesIO(base64.b64decode(response.json()['image'].split(',')[1]))).convert('RGB')
    assert any((lambda p:p[0]>140 and p[1]<80 and p[2]<80)(im.getpixel((x,y))) for y in range(290,350) for x in range(420,580))


def test_thin_lens_static_animation_and_focal_state_survive_pptx():
    diagram = dict(kind='thin_lens', title='凸透镜物距与像距', focal_length_cm=10,
                   object_distance_cm=30, object_height_cm=3,
                   animation=dict(param='object_distance_cm', to=5, duration_ms=2000))
    client = TestClient(app)
    response = client.post('/render/scientific', json=dict(diagram=diagram, frames=True))
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['metadata']['image_distance_cm'] == 15
    assert result['metadata']['magnification'] == -.5
    assert len(result['frames']) == 26
    assert '理想薄透镜' in result['summary']
    focal = client.post('/render/scientific', json=dict(diagram={**diagram, 'object_distance_cm':10, 'animation':None}))
    assert focal.status_code == 200, focal.text
    assert focal.json()['metadata']['image_distance_cm'] is None
    assert '不成有限远像' in focal.json()['summary']
    response = client.post('/render/pptx', json=dict(title='光学',slides=[dict(
        title='移动物体观察成像', scientific=diagram, bullets=['比较实像、焦点状态与虚像。'])]))
    assert response.status_code == 200, response.text
    ppt = Presentation(io.BytesIO(response.content))
    assert '理想薄透镜' in ppt.slides[0].notes_slide.notes_text_frame.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as z:
        media = [n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.gif')]
        assert len(media) == 1
        with Image.open(io.BytesIO(z.read(media[0]))) as image:
            assert image.n_frames >= 25
    from test_render import assert_inside_slide
    assert_inside_slide(ppt)


def test_molecular_structure_highlights_real_functional_group_in_pptx():
    diagram=dict(kind='molecule',title='乙醇中的羟基',smiles='CCO',highlight_group='hydroxyl',
                 animation=dict(param='highlight',to=1,duration_ms=1500))
    response=TestClient(app).post('/render/scientific',json=dict(diagram=diagram,frames=True))
    assert response.status_code==200,response.text
    data=response.json()
    assert data['metadata']['formula']=='C2H6O'
    assert data['metadata']['highlight_atoms']==[2]
    assert data['frames'][0]!=data['frames'][-1]
    response=TestClient(app).post('/render/pptx',json=dict(title='化学结构',slides=[dict(title='识别羟基',scientific=diagram)]))
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content))
    assert any(s.name=='scientific-molecule' for s in ppt.slides[0].shapes)
    assert '不表示反应已验证' in ppt.slides[0].notes_slide.notes_text_frame.text


def test_invalid_structure_and_absent_functional_group_are_rejected():
    client=TestClient(app)
    for change in [dict(smiles='not a molecule'),dict(smiles='CC',highlight_group='hydroxyl')]:
        response=client.post('/render/pptx',json=dict(title='结构',slides=[dict(title='结构核对',scientific={
            'kind':'molecule','title':'结构核对','smiles':'CCO','highlight_group':'none',**change})]))
        assert response.status_code==422,response.text


def test_cell_layers_highlight_only_selected_structures_and_export():
    diagram=dict(kind='cell',title='植物细胞结构',cell_type='plant',
                 visible_parts=['cell_wall','cell_membrane','cytoplasm','nucleus','chloroplast','vacuole'],
                 highlight_parts=['chloroplast'],animation=dict(param='highlight',to=1,duration_ms=1500))
    client=TestClient(app)
    response=client.post('/render/scientific',json=dict(diagram=diagram,frames=True))
    assert response.status_code==200,response.text
    data=response.json()
    assert data['metadata']['highlight_parts']==['chloroplast']
    assert data['frames'][0]!=data['frames'][-1]
    response=client.post('/render/pptx',json=dict(title='细胞',slides=[dict(title='辨认叶绿体',scientific=diagram)]))
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content))
    assert any(s.name=='scientific-cell' for s in ppt.slides[0].shapes)
    assert '不代表真实比例' in ppt.slides[0].notes_slide.notes_text_frame.text
    response=client.post('/render/pptx',json=dict(title='细胞',slides=[dict(title='错误结构',scientific={**diagram,'cell_type':'animal'})]))
    assert response.status_code==422


def test_composed_shapes_paths_labels_and_motion_become_real_ppt_gif():
    diagram=dict(kind='scene',title='颗粒沿路径移动（教学示意）',scene=dict(duration_ms=3000,
        objects=[dict(id='container',type='rect',x=.4,y=.4,width=.6,height=.5,color='#1767bf'),
                 dict(id='route',type='path',points=[[.2,.5],[.4,.3],[.6,.5]],color='#60758a'),
                 dict(id='particle',type='ellipse',x=.2,y=.5,width=.04,height=.06,color='#bc582b'),
                 dict(id='label',type='text',x=.5,y=.8,text='结构化路径演示，非物理仿真',color='#18324c')],
        actions=[dict(object='particle',op='follow_path',path='route',start_ms=0,duration_ms=2500),
                 dict(object='label',op='show',start_ms=1500,duration_ms=500)]))
    client=TestClient(app)
    result=client.post('/render/scientific',json=dict(diagram=diagram,frames=True))
    assert result.status_code==200,result.text
    bundle=result.json();assert len(bundle['frames'])==26
    assert bundle['frames'][0]!=bundle['frames'][-1]
    ppt=client.post('/render/pptx',json=dict(title='通用图形动画',slides=[dict(title='观察运动路径',scientific=diagram)]))
    assert ppt.status_code==200,ppt.text
    with zipfile.ZipFile(io.BytesIO(ppt.content)) as z:
        names=[n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.gif')]
        assert len(names)==1
        with Image.open(io.BytesIO(z.read(names[0]))) as im:assert im.n_frames>=20
    malicious={**diagram,'scene':{**diagram['scene'],'objects':[dict(id='code',type='script',code='alert(1)')]}}
    assert client.post('/render/scientific',json=dict(diagram=malicious)).status_code==422


def test_scene_rejects_invalid_tracks_bounds_and_code_at_render_boundary():
    import copy
    diagram=dict(kind='scene',title='受控路径',scene=dict(duration_ms=3000,
        objects=[dict(id='route',type='path',points=[[.2,.5],[.8,.5]]),
                 dict(id='dot',type='ellipse',x=.2,y=.5,width=.04,height=.06)],
        actions=[dict(object='dot',op='follow_path',path='route',duration_ms=2500)]))
    client=TestClient(app)
    response=client.post('/render/scientific',json=dict(diagram=diagram))
    assert response.status_code==200,response.text
    cases=[]
    changed=copy.deepcopy(diagram)
    changed['scene']['actions'].append(dict(object='route',op='move',x=.4,y=.5,duration_ms=1000))
    cases.append(changed)
    changed=copy.deepcopy(diagram);changed['scene']['actions'][0]['path']='missing';cases.append(changed)
    changed=copy.deepcopy(diagram)
    changed['scene']['actions'].append(dict(object='dot',op='move',x=.5,y=.5,duration_ms=1000))
    cases.append(changed)
    changed=copy.deepcopy(diagram);changed['scene']['objects'][1]['script']='alert(1)';cases.append(changed)
    changed=copy.deepcopy(diagram);changed['scene']['objects'][0]['points'][1][0]=1;cases.append(changed)
    for bad in cases:
        response=client.post('/render/scientific',json=dict(diagram=bad))
        assert response.status_code==422,response.text
        response=client.post('/render/pptx',json=dict(title='校验',slides=[dict(title='输入校验',scientific=bad)]))
        assert response.status_code==422,response.text


def test_static_legacy_tool_rejects_unimplemented_animation():
    response=TestClient(app).post('/render/pptx',json=dict(title='校验',slides=[dict(title='三角形',scientific=dict(
        kind='triangle',title='三角形',sides=[3,4,5],animation=dict(param='highlight',to=1,duration_ms=1500)))]))
    assert response.status_code==422,'Unsupported motion must not silently become a static picture'
