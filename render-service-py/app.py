"""Pure parsing and PPTX rendering. Generated Python is never executed here."""
import io
import base64
import logging
import math
from fractions import Fraction
import unicodedata
import os
import re
import zipfile
from pathlib import Path
from typing import Literal
from teaching_plots import add_plot
from fastapi import FastAPI, HTTPException, UploadFile, Request, Query
from fastapi.responses import Response, JSONResponse
from fastapi.exceptions import RequestValidationError
from pydantic import BaseModel, Field, FiniteFloat, model_validator
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

app=FastAPI(title='Education render service')


class RenderError(ValueError):
    def __init__(self, code, field='content', page=None):
        self.code,self.field,self.page=code,field,page
        super().__init__(code)


@app.exception_handler(RequestValidationError)
async def validation_error(request, exc):
    # Do not echo Pydantic's input, which can contain private teaching materials.
    loc=exc.errors()[0].get('loc',())
    page=loc[2]+1 if len(loc)>2 and loc[1]=='slides' and isinstance(loc[2],int) else None
    field=loc[3] if page and len(loc)>3 else 'content'
    return JSONResponse(status_code=422,content={'detail':{'code':'content_invalid','page':page,'field':field}})
# background, text, accent, surface, soft accent, secondary text
THEMES={
    'classroom':('FFFFFF','182820','009B51','FFFFFF','FFF2B3','476355'),
    'forest':('F5F6EF','233E33','38715B','FFFFFF','E2EADC','586D60'),
    'violet':('F6F3FC','33244D','7651BC','FFFFFF','EAE1F7','716080'),
    'ocean':('FFFFFF','15345B','164B9B','F5F8FC','E8EFF8','61718A'),
    'midnight':('101D30','EEF5FF','69DACC','1A2C43','233E53','AEC3D9'),
    'paper':('F5F0E5','352F28','9B4A31','FCF9F2','EAE1D1','736658'),
    'coral':('FFF5EE','593A34','B84F3F','FFFCF7','FAE0D2','886257'),
}
FONT='Noto Sans CJK SC' if os.name!='nt' else 'Microsoft YaHei'
_MATHTEXT_FONT=None


def mathtext_font():
    """Face used for local raster formula rendering, resolved once.

    FONT above is a declaration written into the .pptx for the *viewer's*
    machine and need not exist on this host. matplotlib draws any glyph the
    chosen face lacks as tofu and only warns, so resolve against what is
    actually installed instead of trusting the platform assumption.
    """
    global _MATHTEXT_FONT
    if _MATHTEXT_FONT is None:
        from matplotlib import font_manager
        installed={f.name for f in font_manager.fontManager.ttflist}
        for name in ('Noto Sans CJK SC','Source Han Sans SC','Microsoft YaHei',
                     'SimHei','WenQuanYi Zen Hei','Noto Sans CJK JP'):
            if name in installed:
                _MATHTEXT_FONT=name
                break
        else:
            _MATHTEXT_FONT='DejaVu Sans'
            logging.getLogger('render').warning(
                'no CJK font installed; Chinese formula glyphs render as boxes. '
                'On Debian/Ubuntu install fonts-noto-cjk.')
    return _MATHTEXT_FONT


LAYOUTS = Literal['cover','concept','steps','comparison','example','exercise','summary']
LAYOUT_LABELS = {'cover':'课程导入','concept':'核心概念','steps':'探究步骤',
                 'comparison':'对照分析','example':'例题解析','exercise':'课堂练习','summary':'学习小结'}

class Chart(BaseModel):
    labels:list[str]=Field(min_length=1,max_length=8)
    values:list[FiniteFloat]=Field(min_length=1,max_length=8)
    title:str=Field(default='实验数据',max_length=80)
    style:Literal['bar','line','scatter','errorbar']='bar'
    x_values:list[FiniteFloat]=Field(default_factory=list,max_length=8)
    y_errors:list[FiniteFloat]=Field(default_factory=list,max_length=8)
    x_label:str=Field(default='',max_length=40)
    y_label:str=Field(default='',max_length=40)
    source_label:str=Field(default='示意数据，非实测',max_length=100)

    @model_validator(mode='after')
    def check_numeric_chart(self):
        if self.style in ('scatter','errorbar'):
            if len(self.labels)!=len(self.values) or any(not label.strip() or len(label)>18 or '\n' in label for label in self.labels):raise ValueError('Invalid chart labels')
            if len(self.x_values)!=len(self.values) or not self.x_label.strip() or not self.y_label.strip() or not self.source_label.strip():raise ValueError('Numeric plot needs axes and source')
            if any(abs(v)>1000000 for v in self.x_values+self.values+self.y_errors):raise ValueError('Plot magnitude exceeds limit')
            if self.style=='errorbar' and len(self.y_errors)!=len(self.values):raise ValueError('Missing error sizes')
            if self.y_errors and (len(self.y_errors)!=len(self.values) or any(v<0 for v in self.y_errors) or self.style!='errorbar'):raise ValueError('Invalid error sizes')
        elif self.x_values or self.y_errors:raise ValueError('Numeric spacing requires numeric plot')
        return self

class FunctionSeries(BaseModel):
    kind:Literal['linear','quadratic','inverse']='linear'
    a:FiniteFloat=Field(default=0,ge=-1000,le=1000)
    c:FiniteFloat=Field(default=0,ge=-1000,le=1000)
    k:FiniteFloat=Field(default=0,ge=-1000,le=1000)
    b:FiniteFloat=Field(default=0,ge=-1000,le=1000)
    mark_x:list[FiniteFloat]=Field(default_factory=list,max_length=2)

    def value(self,x):
        if self.kind=='quadratic':return self.a*x*x+self.b*x+self.c
        if self.kind=='inverse':return self.k/x if x else None
        return self.k*x+self.b

class FunctionGraph(BaseModel):
    renderer:Literal['native','matplotlib']='native'
    title:str=Field(min_length=1,max_length=60)
    x_min:FiniteFloat=Field(ge=-1000,le=0)
    x_max:FiniteFloat=Field(ge=0,le=1000)
    y_min:FiniteFloat=Field(ge=-1000,le=0)
    y_max:FiniteFloat=Field(ge=0,le=1000)
    series:list[FunctionSeries]=Field(min_length=1,max_length=4)

    @model_validator(mode='after')
    def check_graph(self):
        if not self.title.strip() or self.x_max-self.x_min<1 or self.y_max-self.y_min<1:
            raise ValueError('Invalid graph axes')
        for series in self.series:
            if series.kind=='quadratic' and not series.a or series.kind=='inverse' and not series.k:raise ValueError('Zero leading coefficient')
            if series.kind=='linear':
                ys=[series.value(x) for x in (self.x_min,self.x_max)]
                visible=max(ys)>=self.y_min and min(ys)<=self.y_max
            else:
                ys=[series.value(self.x_min+(self.x_max-self.x_min)*i/400) for i in range(401)]
                visible=any(y is not None and self.y_min<=y<=self.y_max for y in ys)
            if not visible:raise ValueError('Function outside axes')
            for x in series.mark_x:
                y=series.value(x)
                if y is None or not self.x_min<=x<=self.x_max or not self.y_min<=y<=self.y_max:
                    raise ValueError('Marked point outside axes')
        return self
class TeachingTable(BaseModel):
    title:str=Field(default='',max_length=60)
    headers:list[str]=Field(min_length=2,max_length=4)
    rows:list[list[str]]=Field(min_length=1,max_length=5)

    @model_validator(mode='after')
    def check_cells(self):
        if any(not h.strip() or len(h)>24 for h in self.headers):
            raise ValueError('Invalid table header')
        if any(len(row)!=len(self.headers) or any(not c.strip() or len(c)>60 for c in row) for row in self.rows):
            raise ValueError('Invalid table row')
        return self

class FigureItem(BaseModel):
    label:str=Field(min_length=1,max_length=24)
    detail:str=Field(default='',max_length=60)

    @model_validator(mode='after')
    def check_label(self):
        if not self.label.strip():raise ValueError('Empty figure label')
        return self

class TeachingFigure(BaseModel):
    kind:Literal['flow','timeline','route','interval']='flow'
    title:str=Field(default='',max_length=60)
    items:list[FigureItem]=Field(min_length=2,max_length=4)
    unit:str=Field(default='',max_length=12)
    values:list[FiniteFloat]=Field(default_factory=list,max_length=4)

    @model_validator(mode='after')
    def check_interval(self):
        if self.kind=='interval':
            if not self.unit.strip() or len(self.values)!=len(self.items):
                raise ValueError('Interval needs unit and one value per node')
            if any(b<=a for a,b in zip(self.values,self.values[1:])) or not math.isfinite(self.values[-1]-self.values[0]):
                raise ValueError('Interval positions must increase with finite span')
        return self

class TeachingPicture(BaseModel):
    data_url:str=Field(max_length=180000)
    caption:str=Field(min_length=1,max_length=60)
    source:str=Field(min_length=1,max_length=300)

    @model_validator(mode='after')
    def check_image(self):
        from PIL import Image
        if not self.caption.strip() or not self.source.strip():raise ValueError('Missing attribution')
        try:
            header,encoded=self.data_url.split(',',1)
            formats={'data:image/png;base64':'PNG','data:image/jpeg;base64':'JPEG'}
            if header not in formats:raise ValueError('Unsupported image')
            raw=base64.b64decode(encoded,validate=True)
            if len(raw)>128000:raise ValueError('Image too large')
            with Image.open(io.BytesIO(raw)) as im:
                if im.format!=formats[header] or min(im.size)<32 or max(im.size)>1600:raise ValueError('Invalid image dimensions')
                im.verify()
        except Exception as e:
            raise ValueError('Invalid image') from e
        return self

class ForceVector(BaseModel):
    label:str=Field(min_length=1,max_length=20)
    x:FiniteFloat=Field(ge=-1000,le=1000)
    y:FiniteFloat=Field(ge=-1000,le=1000)

class ScientificDiagram(BaseModel):
    kind:Literal['triangle','forces','circuit']
    title:str=Field(min_length=1,max_length=60)
    sides:list[FiniteFloat]=Field(default_factory=list,max_length=3)
    object:str=Field(default='',max_length=24)
    unit:str=''
    forces:list[ForceVector]=Field(default_factory=list,max_length=4)
    equilibrium:bool=False
    voltage:FiniteFloat=0
    resistances:list[FiniteFloat]=Field(default_factory=list,max_length=3)
    connection:Literal['','series','parallel']=''
    closed:bool=True

    @model_validator(mode='after')
    def check_science(self):
        if not self.title.strip():raise ValueError('Missing scientific purpose')
        if self.kind=='triangle':
            if len(self.sides)!=3 or any(not 0<n<=1000 for n in self.sides):raise ValueError('Invalid side')
            a,b,c=self.sides
            if min(a+b-c,a+c-b,b+c-a)<=0:raise ValueError('Impossible triangle')
        elif self.kind=='forces':
            if not self.object.strip() or self.unit!='N' or not self.forces:raise ValueError('Missing force object')
            if any(not f.label.strip() or math.hypot(f.x,f.y)<1e-6 for f in self.forces):raise ValueError('Empty force')
            total=sum(math.hypot(f.x,f.y) for f in self.forces)
            if self.equilibrium and math.hypot(sum(f.x for f in self.forces),sum(f.y for f in self.forces))>1e-6*max(1,total):raise ValueError('False equilibrium')
        elif not 0<self.voltage<=1000 or self.connection not in ('series','parallel') or not self.resistances or any(not .01<=r<=1000 for r in self.resistances):raise ValueError('Invalid circuit')
        return self

class CanvasOffset(BaseModel):
    dx:FiniteFloat=Field(ge=-.8,le=.8)
    dy:FiniteFloat=Field(ge=-.8,le=.8)
    scale:FiniteFloat=Field(default=1,ge=.8,le=1.5)
    rotation:FiniteFloat=Field(default=0,ge=-180,le=180)
    locked:bool=False


class CanvasObject(BaseModel):
    id:str=Field(pattern=r'^object-[A-Za-z0-9-]{1,64}$')
    kind:Literal['text','image','shape']
    x:FiniteFloat=Field(ge=0,le=1)
    y:FiniteFloat=Field(ge=0,le=1)
    width:FiniteFloat=Field(ge=.01,le=1)
    height:FiniteFloat=Field(ge=.005,le=1)
    rotation:FiniteFloat=Field(default=0,ge=-180,le=180)
    locked:bool=False
    text:str=Field(default='',max_length=600)
    font_size:FiniteFloat=Field(default=24,ge=16,le=60)
    bold:bool=False
    align:Literal['left','center','right']='left'
    color:str=Field(default='182820',pattern=r'^[A-Fa-f0-9]{6}$')
    background:str=Field(default='',pattern=r'^(?:[A-Fa-f0-9]{6})?$')
    shape:Literal['rect','ellipse','line','arrow']='rect'
    line_width:FiniteFloat=Field(default=2,ge=.5,le=8)
    data_url:str=Field(default='',max_length=180000)
    caption:str=Field(default='',max_length=60)

    @model_validator(mode='after')
    def check_object(self):
        if self.x+self.width>1.000001 or self.y+self.height>1.000001:raise ValueError('Object outside page')
        if self.kind=='text' and not self.text.strip():raise ValueError('Empty text object')
        if self.kind=='image':TeachingPicture(data_url=self.data_url,caption=self.caption or '教师配图',source='教师添加')
        elif self.data_url:raise ValueError('Image payload on non-image object')
        return self


def object_region(slide,region,draw,*args,**kwargs):
    start=len(slide.shapes)
    result=draw(slide,*args,**kwargs)
    for item in list(slide.shapes)[start:]:
        node=item._element.xpath('.//p:cNvPr')[0]
        node.set('descr','edu-region='+region+'|'+node.get('descr',''))
    return result


def region_shapes(slide):
    groups={}
    for item in slide.shapes:
        nodes=item._element.xpath('.//p:cNvPr')
        description=nodes[0].get('descr','') if nodes else ''
        if description.startswith('edu-region='):
            key=description.split('|',1)[0].split('=',1)[1]
            groups.setdefault(key,[]).append(item)
    return groups


def canvas_style_values(items):
    # OOXML absolute lengths only. Relative spacing, chart data, connector flips,
    # picture crop and geometry adjustments must retain their original values.
    drawing='http://schemas.openxmlformats.org/drawingml/2006/main'
    word='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    for item in items:
        roots=[item._element]
        if item.has_chart:roots.append(item.chart._chartSpace)
        for root in roots:
            for node in root.iter():
                namespace,_,tag=node.tag[1:].partition('}')
                if namespace==drawing:
                    attrs={'rPr':('sz',),'defRPr':('sz',),'endParaRPr':('sz',),
                           'bodyPr':('lIns','rIns','tIns','bIns'),
                           'tcPr':('marL','marR','marT','marB'),
                           'pPr':('marL','marR','indent'),
                           'ln':('w',),'spcPts':('val',),
                           'gridCol':('w',),'tr':('h',)}.get(tag,())
                    for attr in attrs:
                        value=node.get(attr)
                        if value is not None:
                            yield node,attr,int(value),100 if attr=='sz' else None
                elif namespace==word and tag in ('sz','szCs'):
                    attr='{'+word+'}val'
                    if node.get(attr) is not None:yield node,attr,int(node.get(attr)),2


def canvas_min_scale(items,key):
    floor=10 if key=='visual' else 16
    return max([1 if key=='formula' and any(i.shape_type==13 for i in items) else 0]+[min(floor,value/unit)/(value/unit)
                     for _,_,value,unit in canvas_style_values(items) if unit and value>0])


def canvas_bounds(left,top,width,height,rotation):
    angle=math.radians(rotation)
    w=abs(width*math.cos(angle))+abs(height*math.sin(angle))
    h=abs(width*math.sin(angle))+abs(height*math.cos(angle))
    return left+(width-w)/2,top+(height-h)/2,w,h


def apply_canvas_positions(slide,positions,width,height,data=None):
    groups=region_shapes(slide)
    for key,offset in positions.items():
        if key not in groups:raise RenderError('canvas_object_missing','canvas_positions')
        items=groups[key];scale=offset.scale
        rotatable=key!='formula' and not(key=='visual' and data and (data.graph or data.scientific or data.chart or data.diagram))
        if offset.rotation and not rotatable:raise RenderError('canvas_rotation_protected','canvas_positions')
        min_scale=canvas_min_scale(items,key)
        if scale+1e-6<min_scale:raise RenderError('canvas_text_too_small','canvas_positions')
        x=min(i.left for i in items);y=min(i.top for i in items)
        w=max(i.left+i.width for i in items)-x;h=max(i.top+i.height for i in items)-y
        dx,dy=round(offset.dx*width),round(offset.dy*height)
        cx=x+w*scale/2+dx;cy=y+h*scale/2+dy
        angle=math.radians(offset.rotation)
        transforms=[(item,round(x+(item.left-x)*scale)+dx,round(y+(item.top-y)*scale)+dy,
                     round(item.width*scale),round(item.height*scale)) for item in items]
        rotated=[]
        for item,left,top,iw,ih in transforms:
            rx=left+iw/2-cx;ry=top+ih/2-cy
            left=round(cx+rx*math.cos(angle)-ry*math.sin(angle)-iw/2)
            top=round(cy+rx*math.sin(angle)+ry*math.cos(angle)-ih/2)
            bounds=canvas_bounds(left,top,iw,ih,item.rotation+offset.rotation)
            if min(bounds[0],bounds[1])<-.5 or bounds[0]+bounds[2]>width+.5 or bounds[1]+bounds[3]>height+.5:
                raise RenderError('canvas_outside_slide','canvas_positions')
            rotated.append((item,left,top,iw,ih))
        if scale!=1:
            for node,attr,value,_ in canvas_style_values(items):node.set(attr,str(round(value*scale)))
        for item,left,top,iw,ih in rotated:
            item.left=left;item.top=top;item.width=iw;item.height=ih;item.rotation=(item.rotation+offset.rotation)%360
        node=items[0]._element.xpath('.//p:cNvPr')[0]
        node.set('descr',node.get('descr','')+f'|edu-base-min={min_scale}|edu-scale={scale}|edu-box={x+dx},{y+dy},{round(w*scale)},{round(h*scale)}|edu-rotation={offset.rotation}')


def draw_canvas_object(slide,obj,width,height):
    x,y,w,h=obj.x*width,obj.y*height,obj.width*width,obj.height*height
    bounds=canvas_bounds(x,y,w,h,obj.rotation)
    if min(bounds[0],bounds[1])<-.5 or bounds[0]+bounds[2]>width+.5 or bounds[1]+bounds[3]>height+.5:
        raise RenderError('canvas_outside_slide','canvas_objects')
    if obj.kind=='text':
        # Explicit type size is preserved; text must fit instead of silently shrinking.
        item=box(slide,x/914400,y/914400,w/914400,h/914400,obj.text,round(obj.font_size),obj.color,obj.bold)
        for p in item.text_frame.paragraphs:p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[obj.align]
        if obj.background:item.fill.solid();item.fill.fore_color.rgb=rgb(obj.background)
    elif obj.kind=='image':
        item=slide.shapes.add_picture(io.BytesIO(base64.b64decode(obj.data_url.split(',',1)[1],validate=True)),round(x),round(y),width=round(w),height=round(h))
    elif obj.shape=='line':
        item=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,round(x),round(y+h/2),round(x+w),round(y+h/2))
        # Connector geometry remains horizontal; metadata retains the hit box.
        item.line.color.rgb=rgb(obj.color);item.line.width=Pt(obj.line_width)
    else:
        item=shape(slide,x/914400,y/914400,w/914400,h/914400,obj.background or 'FFFFFF',{'rect':MSO_SHAPE.RECTANGLE,'ellipse':MSO_SHAPE.OVAL,'arrow':MSO_SHAPE.RIGHT_ARROW}[obj.shape],obj.color)
        if not obj.background:item.fill.background()
        item.line.width=Pt(obj.line_width)
    item.rotation=obj.rotation%360
    node=item._element.xpath('.//p:cNvPr')[0]
    node.set('descr',f'edu-region={obj.id}|edu-rotatable=1|edu-box={round(x)},{round(y)},{round(w)},{round(h)}|edu-rotation={obj.rotation}|'+obj.caption.replace('|','／'))


def apply_canvas_order(slide,order):
    groups=region_shapes(slide)
    if any(key not in groups for key in order):raise RenderError('canvas_object_missing','canvas_order')
    # Each native teaching region remains contiguous; ordering cannot separate axes/labels.
    for key in [key for key in groups if key not in order]+order:
        for item in groups[key]:slide.shapes._spTree.remove(item._element);slide.shapes._spTree.insert_element_before(item._element,'p:extLst')


class Slide(BaseModel):
    canvas_objects:list[CanvasObject]=Field(default_factory=list,max_length=24)
    canvas_order:list[str]=Field(default_factory=list,max_length=29)
    canvas_positions:dict[Literal['title','body','visual','formula','focus'],CanvasOffset]=Field(default_factory=dict,max_length=5)
    answer:str=Field(default='',max_length=2000)
    scientific:ScientificDiagram|None=None
    graph:FunctionGraph|None=None
    title:str=Field(min_length=1,max_length=70)
    bullets:list[str]=Field(default_factory=list,max_length=6)
    notes:str=Field(default='',max_length=12000)
    formula:str=Field(default='',max_length=250)
    chart:Chart|None=None
    diagram:Literal['','buoyancy','triangle']=''
    layout:LAYOUTS|None=None
    section:str=Field(default='',max_length=80)
    focus:str=Field(default='',max_length=60)
    table:TeachingTable|None=None
    figure:TeachingFigure|None=None
    picture:TeachingPicture|None=None
    emphasis:list[str]=Field(default_factory=list,max_length=6)

    @model_validator(mode='after')
    def check_media(self):
        ids=[obj.id for obj in self.canvas_objects]
        if len(ids)!=len(set(ids)) or len(self.canvas_order)!=len(set(self.canvas_order)):raise ValueError('Duplicate canvas IDs')
        if self.scientific and (self.graph or self.picture or self.chart or self.table or self.figure or self.diagram):raise ValueError('Scientific diagram requires own panel')
        if self.graph and (self.picture or self.chart or self.table or self.figure or self.diagram):
            raise ValueError('Function graph requires its own panel')
        if any(not term.strip() or len(term)>20 for term in self.emphasis):raise ValueError('Invalid emphasis')
        if self.picture and (self.chart or self.diagram or (self.table and self.figure)):
            raise ValueError('Too many media objects')
        if (self.table or self.figure) and (self.chart or self.diagram):
            raise ValueError('Split conflicting media onto separate slides')
        return self
class Deck(BaseModel):
    editable_formulas:bool=False
    title:str=Field(min_length=1,max_length=100)
    theme:Literal['forest','violet','ocean','midnight','paper','coral','classroom']='forest'
    slides:list[Slide]=Field(min_length=1,max_length=100)
    needs_review:bool=False

    @model_validator(mode='after')
    def check_total_images(self):
        if sum((len(s.picture.data_url) if s.picture else 0)+sum(len(o.data_url) for o in s.canvas_objects) for s in self.slides)>1200000:
            raise ValueError('Total image budget exceeded')
        return self

def rgb(h):return RGBColor.from_string(h)
def inline_font_content(text):
    # Plain body text supports no TeX styling. Unwrap only font commands,
    # preserving the actual characters, operators and grouping inside them.
    while True:
        plain=re.sub(r'\\+(?:mathbf|mathrm|mathit|textbf|textit|text)\{([^{}]*)\}',lambda m:m[1],text)
        if plain==text:return text
        text=plain

def keep_number_unit(text):
    # Preserve stored content; keep the value and unit together in display text.
    return re.sub(r'(\d(?:\.\d+)?)[ \u00a0]+((?:km/h|m/s[²2]?|mm|cm|km|ms|min|kg|mL|mol|Hz|Pa|s|h|m[²³]?|g|L|N|J|W|K|V|A|°C|℃|%)(?![A-Za-z]))',lambda m:m[1]+'\u00a0'+m[2],text)

def text_lines(text, width, size):
    """Conservative word wrapping, including explicit breaks and unbroken tokens."""
    capacity = max(1, (width * 72 - 8) / size - .7)
    count = 0
    for line in keep_number_unit(inline_font_content(text)).expandtabs(4).split('\n'):
        used = 0
        count += 1
        for token in re.findall(r'\d+(?:\.\d+)?\u00a0[A-Za-z/²³%℃°0-9]+|[A-Za-z0-9_]+|.', line):
            units = sum(1.05 if unicodedata.east_asian_width(c) in 'WF' or c in '“”‘’—…' else
                        .9 if c in 'MW@%&' else .62 for c in token)
            if used and used + units > capacity:
                count += 1
                used = 0
            while units > capacity:
                count += 1
                units -= capacity
            used += units
    return count


def box(slide,x,y,w,h,text,size,color,bold=False,minimum=None,font=FONT):
    # ponytail: conservative font metrics, not a shaping engine. Reject at the
    # readable floor; native PowerPoint preview is the visual acceptance check.
    minimum = size if minimum is None else minimum
    text=keep_number_unit(inline_font_content(text).replace('\r\n','\n').replace('\r','\n').replace('\v','\n').expandtabs(4))
    size = next((pt for pt in range(size, minimum - 1, -1)
                 if text_lines(text,w,pt)*pt*1.4 <= h*72-3),None)
    if size is None:
        raise RenderError('text_overflow')
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    frame=shape.text_frame;frame.word_wrap=True;frame.auto_size=MSO_AUTO_SIZE.NONE
    frame.margin_left=frame.margin_top=frame.margin_bottom=0
    # PowerPoint can render CJK glyphs slightly past its wrap boundary. The
    # measured inset is also subtracted by text_lines, not hidden clipping.
    frame.margin_right=Pt(8)
    for i,line in enumerate(text.split('\n')):
        p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.font._element.set('lang','zh-CN')
        p.font._element.append(parse_xml(f'<a:ea {nsdecls("a")} typeface="{font}"/>'))
        p._p.get_or_add_pPr().set('hangingPunct','0')
        p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=1.15
    return shape


def shape(slide,x,y,w,h,color,kind=MSO_SHAPE.RECTANGLE,line=None):
    sh=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(color)
    if line:sh.line.color.rgb=rgb(line)
    else:sh.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:sh.adjustments[0]=.12
    sh._element.spPr.append(parse_xml(f'<a:effectLst {nsdecls("a")}/>'))
    return sh


def choose_layout(data,index,total):
    if data.layout:return data.layout
    for words,layout in [('练习|练一练|检测','exercise'),('例题|示例','example'),
                         ('小结|总结|回顾','summary'),('比较|对比|辨析','comparison'),
                         ('步骤|流程|实验|探究','steps')]:
        if re.search(words,data.title):return layout
    if index==0 and len(data.bullets)<=3 and sum(map(len,data.bullets))<=100:return 'cover'
    if index==total-1 and total>1:return 'summary'
    return 'concept'


def decorate(slide,theme,cover):
    bg,ink,accent,surface,soft,muted=THEMES[theme]
    if theme=='classroom':
        shape(slide,.48,.45,.055,.8,accent)
        shape(slide,11.8,0,1.53,.18,'FFD539')
        shape(slide,12.99,.18,.343,1.02,'FFD539')
        shape(slide,0,7.1,.42,.4,'FFD539')
        if cover:
            shape(slide,9.05,1.6,3.1,4.8,soft,MSO_SHAPE.ROUNDED_RECTANGLE)
            for i in range(3):
                shape(slide,9.5,2.35+i,2.15,.12,accent)
                shape(slide,9.3,2.25+i,.32,.32,'FFD539',MSO_SHAPE.OVAL)
    elif theme=='forest':
        shape(slide,0,0,.18,7.5,accent)
        shape(slide,.7,.87,.65,.055,accent)
        if cover:
            shape(slide,8.9,1.35,3.7,5.3,soft,MSO_SHAPE.ROUNDED_RECTANGLE)
            for x,y,r in [(9.3,2,1.9),(10.5,3.25,1.55),(9.45,4.75,1.15)]:
                shape(slide,x,y,r,r,accent,MSO_SHAPE.OVAL)
    elif theme=='violet':
        shape(slide,0,0,13.333,.12,accent)
        shape(slide,11.75,.39,.85,.34,soft,MSO_SHAPE.ROUNDED_RECTANGLE)
        if cover:
            shape(slide,8.9,1.55,3.55,3.55,soft,MSO_SHAPE.OVAL)
            shape(slide,9.2,3.4,2.8,2.8,accent,MSO_SHAPE.DIAMOND)
            shape(slide,11.65,1.6,.5,.5,accent,MSO_SHAPE.OVAL)
    elif theme=='ocean':
        shape(slide,.55,.95,.045,.4,accent)
        shape(slide,.7,6.96,11.9,.012,'DFE7F2')
        if cover:
            shape(slide,8.65,1.3,3.95,5.42,accent)
            for i in range(8):
                shape(slide,8.9+i*.43,3.6,.012,2.7,'4472B8')
                shape(slide,8.9,3.6+i*.34,3.25,.012,'4472B8')
            connector=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(9),Inches(5.8),Inches(12),Inches(2.2))
            connector.line.color.rgb=rgb('D9E7FF');connector.line.width=Pt(2.2)
    elif theme=='midnight':
        for i in range(7):shape(slide,9+i*.52,.28,.025,.38,soft)
        shape(slide,.7,.89,.5,.035,accent)
        if cover:
            for i in range(6):
                shape(slide,8.8+i*.72,1.7,.015,4.8,soft)
                shape(slide,8.8,1.7+i*.8,3.62,.015,soft)
            shape(slide,9.35,2.45,2.5,2.5,soft,MSO_SHAPE.OVAL,accent)
            shape(slide,11.35,2.2,.4,.4,accent,MSO_SHAPE.OVAL)
    elif theme=='paper':
        shape(slide,.7,.87,11.9,.015,ink)
        shape(slide,.7,6.91,11.9,.012,ink)
        if cover:
            shape(slide,8.75,1.3,.02,5.2,ink)
            box(slide,9.25,2.1,3,2.1,'Aa',96,accent,True,font='Georgia')
            box(slide,9.3,4.6,2.8,.5,'观察 · 思考 · 表达',16,muted)
    else:
        shape(slide,.7,.37,.33,.33,accent,MSO_SHAPE.OVAL)
        shape(slide,1.12,.43,.22,.22,'E9B454',MSO_SHAPE.OVAL)
        if cover:
            shape(slide,9.15,1.8,3.25,3.25,soft,MSO_SHAPE.OVAL)
            shape(slide,8.95,4.35,3.45,1.8,surface,MSO_SHAPE.ROUNDED_RECTANGLE)
            shape(slide,9.4,4.8,2.45,.12,accent)
            shape(slide,9.4,5.2,1.5,.12,'E9B454')


def body(slide,data,layout,theme,x,y,w,h):
    _,ink,accent,surface,soft,muted=THEMES[theme]
    items=data.bullets
    if not items:return
    if len(items)==1 and layout!='cover':
        box(slide,x,y,w,h,items[0],28,ink,True,minimum=18)
        return
    rounded=MSO_SHAPE.RECTANGLE if theme in ('paper','midnight','ocean') else MSO_SHAPE.ROUNDED_RECTANGLE

    def panel(px,py,pw,ph,text,highlight=False,size=23,padding=.17):
        shape(slide,px,py,pw,ph,soft if highlight else surface,rounded)
        if theme=='paper':shape(slide,px,py,.025,ph,accent)
        elif theme=='ocean':shape(slide,px,py,.035,ph,accent)
        box(slide,px+.22,py+padding,pw-.44,ph-2*padding,text,size,ink,highlight,minimum=16)

    if layout=='cover':
        row=h/len(items)
        for i,text in enumerate(items):
            shape(slide,x,y+i*row+.12,.07,.2,accent)
            box(slide,x+.25,y+i*row,w-.25,row-.04,text,23,ink,minimum=16)
    elif layout=='concept':
        # Key idea first, supporting evidence below. Avoid a grid of identical cards.
        lead=max(h*.43,text_lines(items[0],w,22)*22*1.4/72+.15) if len(items)>1 else h
        box(slide,x,y,w,lead-.1,items[0],30,ink,True,minimum=22)
        rest=items[1:]
        if rest:
            cols=2 if len(rest)>1 and w>=6 else 1
            rows=math.ceil(len(rest)/cols)
            cellw=(w-.28*(cols-1))/cols;cellh=(h-lead-.08)/rows
            for i,text in enumerate(rest):
                px=x+(i%cols)*(cellw+.28);py=y+lead+.08+(i//cols)*cellh
                box(slide,px,py,.62,.38,f'{i+1:02}',15,accent,True)
                if cellh<1.15:
                    box(slide,px+.62,py,cellw-.62,cellh-.08,text,22,ink,minimum=16)
                else:
                    box(slide,px,py+.43,cellw,cellh-.48,text,22,ink,minimum=16)
    elif layout=='steps':
        # Editable, open numbered sequence; height follows wrapped text.
        heights=[max(.65,text_lines(t,w-.75,16)*16*1.4/72+.15) for t in items]
        if sum(heights)>h:raise RenderError('text_overflow')
        spare=(h-sum(heights))/len(items)
        heights=[height+spare for height in heights]
        py=y
        for i,(text,row) in enumerate(zip(items,heights)):
            box(slide,x,py,.48,.45,str(i+1),17,accent,True)
            box(slide,x+.7,py,w-.7,row-.1,text,23,ink,minimum=16)
            py+=row
    elif layout=='comparison':
        if w<6:
            # Evidence occupies the other column. Keep A/B comparisons distinct
            # instead of reclassifying their first statement as a key concept.
            row=h/len(items)
            for i,text in enumerate(items):
                py=y+i*row
                shape(slide,x,py,w,.02,accent)
                box(slide,x,py+.06,.36,.32,'A' if i%2==0 else 'B',14,accent,True)
                box(slide,x+.5,py+.06,w-.5,row-.1,text,24,ink,minimum=18)
            return
        col=(w-.26)/2
        for side in range(2):
            px=x+side*(col+.26)
            shape(slide,px,y,col,.035,accent)
            box(slide,px+.22,y+.17,col-.44,.35,'A  /  对照一' if side==0 else 'B  /  对照二',15,accent,True)
            subset=items[side::2]
            row=(h-.85)/max(1,len(subset))
            for i,text in enumerate(subset):
                box(slide,px+.22,y+.77+i*row,col-.44,row-.14,text,23,ink,minimum=16)
    elif layout=='example':
        left=w*.43
        box(slide,x,y,left-.2,h,items[0],27,ink,True,minimum=18)
        rest=items[1:]
        if rest:
            row=h/len(rest)
            for i,text in enumerate(rest):
                py=y+i*row
                box(slide,x+left,py,.48,.4,str(i+1),16,accent,True)
                box(slide,x+left+.62,py,w-left-.62,row-.12,text,23,ink,minimum=16)
    elif layout=='exercise':
        lead=min(h*.45,max(.8,text_lines(items[0],w-.44,22)*22*1.4/72+.34))
        panel(x,y,w,lead,items[0],True,27)
        rest=items[1:]
        if rest:
            row=(h-lead-.22)/len(rest)
            for i,text in enumerate(rest):
                box(slide,x+.25,y+lead+.22+i*row,w-.5,row-.1,text,22,ink,minimum=16)
        else:
            for i in range(3):shape(slide,x+.25,y+lead+.3+i*.35,w-.5,.015,soft)
    else: # summary: one central takeaway and a short supporting list
        if len(items)==1:
            box(slide,x,y,w,h,items[0],30,ink,True,minimum=20)
        else:
            left=w*.45
            box(slide,x,y,left-.15,h,items[0],29,ink,True,minimum=19)
            row=h/(len(items)-1)
            for i,text in enumerate(items[1:]):
                py=y+i*row
                box(slide,x+left,py,.62,.38,f'{i+2:02}',15,accent,True)
                if row<1.15:
                    box(slide,x+left+.62,py,w-left-.62,row-.08,text,22,ink,minimum=16)
                else:
                    box(slide,x+left,py+.43,w-left,row-.48,text,22,ink,minimum=16)


def function_graph(slide,data,theme,x,y,w,h,show_coordinates=True):
    _,ink,accent,surface,soft,muted=THEMES[theme]
    if data.renderer=='matplotlib':
        box(slide,x+.2,y+.05,w-.4,.5,data.title,17,accent,True,minimum=14)
        add_plot(slide,data,THEMES[theme],mathtext_font(),x,y+.6,w,h-.6,True,show_coordinates)
        return
    shape(slide,x,y,w,h,surface if theme=='midnight' else 'FFFFFF',MSO_SHAPE.RECTANGLE,'DFE7F2')
    box(slide,x+.2,y+.12,w-.4,.4,data.title,17,accent,True,minimum=14)
    # The same inch/unit ratio on both axes makes slope comparisons honest.
    legendh=.38*len(data.series)
    left,top=x+.85,y+.72
    availw,availh=w-1.28,h-1.22-legendh
    if min(availw,availh)<.9:raise RenderError('text_overflow')
    scale=min(availw/(data.x_max-data.x_min),availh/(data.y_max-data.y_min))
    pw,ph=(data.x_max-data.x_min)*scale,(data.y_max-data.y_min)*scale
    left+=(availw-pw)/2;top+=(availh-ph)/2
    px=lambda value:left+(value-data.x_min)*scale
    py=lambda value:top+(data.y_max-value)*scale
    def line(name,x1,y1,x2,y2,color,width):
        item=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        item.name=name;item.line.color.rgb=rgb(color);item.line.width=Pt(width)
        return item
    def tick_step(span):
        raw=span/4;power=10**math.floor(math.log10(raw))
        return next(n*power for n in (1,2,5,10) if n*power>=raw)
    for axis,lo,hi in [('x',data.x_min,data.x_max),('y',data.y_min,data.y_max)]:
        step=tick_step(hi-lo)
        while step*scale<(.88 if axis=='x' else .4):step*=2
        value=math.ceil(lo/step)*step
        while value<=hi+step*.0001:
            if axis=='x':
                line('function-grid-x',px(value),top,px(value),top+ph,soft,.6)
                box(slide,px(value)-.425,top+ph+.04,.85,.34,f'{value:g}',14,muted)
            else:
                line('function-grid-y',left,py(value),left+pw,py(value),soft,.6)
                if abs(value)>step*.0001:box(slide,left-.85,py(value)-.17,.72,.34,f'{value:g}',14,muted)
            value+=step
    line('function-axis-x',left,py(0),left+pw,py(0),muted,1.2)
    line('function-axis-y',px(0),top,px(0),top+ph,muted,1.2)
    box(slide,left+pw+.07,py(0)-.17,.3,.34,'x',14,ink)
    box(slide,px(0)+.08,top,.3,.34,'y',14,ink)
    colors=[accent,'D75454','168979','8059B0']
    for i,series in enumerate(data.series):
        color=colors[i]
        if series.kind=='linear':
            xa,xb=data.x_min,data.x_max
            if series.k:
                ya,yb=sorted(((data.y_min-series.b)/series.k,(data.y_max-series.b)/series.k))
                xa,xb=max(xa,ya),min(xb,yb)
            line(f'function-line-{i}',px(xa),py(series.value(xa)),px(xb),py(series.value(xb)),color,2)
        else:
            previous=None
            for n in range(401):
                value=data.x_min+(data.x_max-data.x_min)*n/400;result=series.value(value)
                if result is None or not data.y_min<=result<=data.y_max:previous=None;continue
                if previous and not (series.kind=='inverse' and previous[0]*value<=0):line(f'function-curve-{i}-{n}',px(previous[0]),py(previous[1]),px(value),py(result),color,2)
                previous=(value,result)
        labels=[]
        for j,value in enumerate(series.mark_x):
            result=series.value(value)
            dot=shape(slide,px(value)-.045,py(result)-.045,.09,.09,color,MSO_SHAPE.OVAL)
            dot.name=f'function-point-{i}-{j}'
            def number_label(number):
                fraction=Fraction(number).limit_denominator(12)
                return str(fraction) if abs(float(fraction)-number)<1e-12 else f'{number:g}'
            labels.append(f'({number_label(value)}, {number_label(result)})')
        slope='x' if series.k==1 else '-x' if series.k==-1 else f'{series.k:g}x'
        expression=f'y = {slope}' if series.k else f'y = {series.b:g}'
        if series.k and series.b:expression+=f' {"+" if series.b>0 else "-"} {abs(series.b):g}'
        if series.kind=='quadratic':expression=f'y = {series.a:g}x² {series.b:+g}x {series.c:+g}'
        elif series.kind=='inverse':expression=f'y = {series.k:g}/x'
        text=expression+(' · '+', '.join(labels) if labels and show_coordinates else '')
        box(slide,x+.22,y+h-legendh-.05+i*.38,w-.44,.35,text,14,color,True)


def scientific_diagram(slide,data,theme,x,y,w,h):
    _,ink,accent,_,soft,muted=THEMES[theme]
    box(slide,x+.1,y,w-.2,.5,data.title,20,accent,True,minimum=18)
    def line(name,x1,y1,x2,y2,arrow=False,color=accent):
        sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        sh.name='scientific-'+name;sh.line.color.rgb=rgb(color);sh.line.width=Pt(2)
        if arrow:sh.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle"/>'))
    if data.kind=='triangle':
        a,b,c=data.sides;cx=(a*a+b*b-c*c)/(2*a);cy=math.sqrt(max(0,b*b-cx*cx))
        scale=min((w-1.2)/(max(a,cx)-min(0,cx)),(h-1.8)/cy)
        ox=x+.6-min(0,cx)*scale;oy=y+h-.85
        pts=[(ox,oy),(ox+a*scale,oy),(ox+cx*scale,oy-cy*scale)]
        for i,j,label in [(0,1,f'AB = {a:g}'),(0,2,f'AC = {b:g}'),(1,2,f'BC = {c:g}')]:
            p,q=pts[i],pts[j];line(f'side-{i}-{j}',*p,*q)
        box(slide,x+.1,y+h-.36,w-.2,.32,f'AB = {a:g}   AC = {b:g}   BC = {c:g}',12,ink,minimum=10)
        for i,(px,py) in enumerate(pts):box(slide,max(x,min(x+w-.3,px)),max(y+.55,min(y+h-.3,py)),.3,.28,'ABC'[i],12,accent,True,minimum=10)
    elif data.kind=='forces':
        xmin,xmax=min(0,*[f.x for f in data.forces]),max(0,*[f.x for f in data.forces])
        ymin,ymax=min(0,*[f.y for f in data.forces]),max(0,*[f.y for f in data.forces])
        # Fit the actual vector bounds, retaining one scale for both axes.
        # An asymmetric three-force balance should use the available height.
        scale=min((w-2.8)/(xmax-xmin) if xmax>xmin else float('inf'),(h-1.8)/(ymax-ymin) if ymax>ymin else float('inf'))
        ox=x+w/2-(xmin+xmax)*scale/2;oy=y+.95+ymax*scale
        dot=shape(slide,ox-.1,oy-.1,.2,.2,accent,MSO_SHAPE.OVAL);dot.name='scientific-object'
        box(slide,x+.2,y+.55,w-.4,.4,'研究对象：'+data.object,16,ink,minimum=14)
        labels=[]
        colors=[accent,'B45309','167D75','713BB0']
        for i,f in enumerate(data.forces):
            ex,ey=ox+f.x*scale,oy-f.y*scale;line(f'force-{i}',ox,oy,ex,ey,True,colors[i])
            magnitude=math.hypot(f.x,f.y)
            labeled=re.search(r'([0-9]+(?:\.[0-9]+)?)\s*N\s*$',f.label)
            label=f.label if labeled and math.isclose(float(labeled[1]),magnitude,rel_tol=1e-9) else f'{f.label} {magnitude:g} N'
            labelw=2.7
            lx=max(x,min(x+w-labelw,ex+.12 if i%2==0 else ex-labelw-.12))
            ly=max(y+.95,min(y+h-.85,ey-.2))
            ly=next((candidate for step in [0,1,-1,2,-2,3,-3]
                     if y+.95<=(candidate:=ly+step*.42)<=y+h-.85
                     and not any(abs(lx-px)<labelw and abs(candidate-py)<.419 for px,py in labels)),None)
            if ly is None:raise RenderError('scientific_force_labels_dense')
            labels.append((lx,ly))
            box(slide,lx,ly,labelw,.4,label,18,colors[i],True)
        sx,sy=sum(f.x for f in data.forces),sum(f.y for f in data.forces)
        box(slide,x+.1,y+h-.4,w-.2,.35,f'合力 ({sx:g}, {sy:g}) N'+(' · 平衡' if data.equilibrium else ''),12,muted,minimum=10)
    else:
        left,right,top,bottom=x+.5,x+w-.5,y+1.15,y+h-1.35
        if data.connection=='series':line('wire-bottom',left,bottom,right,bottom)
        line('wire-left',left,top,left,bottom);line('wire-right',right,top,right,bottom)
        mid=(left+right)/2
        if data.closed:line('source-left',left,top,mid-.13,top)
        else:
            line('switch-left',left,top,left+.3,top);line('switch-open',left+.3,top,left+.6,top-.16);line('source-left',left+.7,top,mid-.13,top)
        line('source-right',mid+.13,top,right,top)
        line('source-positive',mid-.13,top-.2,mid-.13,top+.2);line('source-negative',mid+.13,top-.12,mid+.13,top+.12)
        box(slide,x+.2,y+.55,w-.4,.4,f'理想电源 {data.voltage:g} V · '+('闭合' if data.closed else '断开（电流为0）'),18,ink)
        def resistor(i,cx,cy,r):
            sh=shape(slide,cx-.32,cy-.12,.64,.24,'FFFFFF',MSO_SHAPE.RECTANGLE,accent);sh.name=f'scientific-resistor-{i}'
            # Alternate series labels above/below the same wire; component
            # positions and connection remain exact even with three resistors.
            labely=(cy-.6 if i%2 else cy+.12) if data.connection=='series' else cy-.44
            labelx=cx-.9 if data.connection=='series' else left+.15
            box(slide,labelx,labely,1.8,.4,f'R{i+1}={r:g} Ω',18,ink)
        if data.connection=='series':
            for i,r in enumerate(data.resistances):resistor(i,left+(right-left)*(i+1)/(len(data.resistances)+1),bottom,r)
            resistance=sum(data.resistances)
        else:
            for i,r in enumerate(data.resistances):
                cy=top+(bottom-top)*(i+1)/len(data.resistances);line(f'branch-{i}',left,cy,right,cy);resistor(i,mid,cy,r)
                current=data.voltage/r if data.closed else 0
                box(slide,right-2.1,cy-.44,2.1,.4,f'I{i+1}={current:.3g} A',18,muted)
            resistance=1/sum(1/r for r in data.resistances)
        current=data.voltage/resistance if data.closed else 0
        box(slide,x+.1,y+h-.82,w-.2,.8,f'总电阻 {resistance:.3g} Ω\n总电流 {current:.3g} A · 理想电路模型',18,ink)

def visual(slide,data,theme,x,y,w,h):
    _,ink,accent,surface,soft,muted=THEMES[theme]
    if data.graph:
        task=data.title+'。'+'。'.join(data.bullets)
        reading_coordinates=data.layout=='exercise' and re.search(r'(?:读出|求出|求|确定)[^。！？?\n]{0,12}坐标',task)
        function_graph(slide,data.graph,theme,x,y,w,h,show_coordinates=not reading_coordinates)
        return
    if data.scientific:
        scientific_diagram(slide,data.scientific,theme,x,y,w,h)
        return
    shape(slide,x,y,w,h,surface,MSO_SHAPE.ROUNDED_RECTANGLE)
    if data.chart:
        c=data.chart
        if c.style in ('scatter','errorbar'):
            box(slide,x+.2,y+.06,w-.4,.5,c.title,17,ink,True,minimum=14)
            add_plot(slide,c,THEMES[theme],mathtext_font(),x,y+.62,w,h-1.06)
            box(slide,x+.2,y+h-.38,w-.4,.34,c.source_label,12,muted,minimum=10)
            return
        if len(c.labels)!=len(c.values):raise RenderError('chart_invalid')
        if any(len(label)>18 or '\n' in label for label in c.labels):raise RenderError('chart_invalid')
        box(slide,x+.2,y+.15,w-.4,.62,c.title,17,ink,True,minimum=14)
        chartdata=ChartData();chartdata.categories=c.labels;chartdata.add_series(c.title,c.values)
        chart_type=XL_CHART_TYPE.LINE_MARKERS if c.style=='line' else XL_CHART_TYPE.COLUMN_CLUSTERED
        ch=slide.shapes.add_chart(chart_type,Inches(x+.15),Inches(y+.85),Inches(w-.3),Inches(h-1),chartdata).chart
        ch.has_legend=False;ch.has_title=False
        ch.font.name=FONT;ch.font.size=Pt(11);ch.font.color.rgb=rgb(ink)
        # python-pptx exposes series/axis fills, but not the chart-area fill.
        area=parse_xml(f'<c:spPr {nsdecls("c","a")}><a:solidFill><a:srgbClr val="{surface}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>')
        ch._chartSpace.insert_element_before(area,'c:txPr','c:externalData','c:printSettings')
        if c.style=='line':
            ch.series[0].format.line.color.rgb=rgb(accent)
            ch.series[0].format.line.width=Pt(2.5)
        else:
            ch.series[0].format.fill.solid();ch.series[0].format.fill.fore_color.rgb=rgb(accent)
            ch.series[0].format.line.fill.background()
        for axis in (ch.category_axis,ch.value_axis):
            axis.tick_labels.font.name=FONT;axis.tick_labels.font.size=Pt(10);axis.tick_labels.font.color.rgb=rgb(muted)
        ch.category_axis.tick_labels.offset=100
        ch.value_axis.has_major_gridlines=True
        ch.value_axis.major_gridlines.format.line.color.rgb=rgb(soft)
    elif data.diagram=='buoyancy':
        shape(slide,x+w*.12,y+h*.4,w*.76,h*.45,soft)
        shape(slide,x+w*.36,y+h*.34,w*.28,h*.3,'D5AD7D',MSO_SHAPE.RECTANGLE,accent)
        shape(slide,x+w*.45,y+h*.14,w*.1,h*.25,accent,MSO_SHAPE.UP_ARROW)
        box(slide,x+w*.59,y+h*.14,w*.32,h*.18,'F浮',20,ink,minimum=16)
        box(slide,x+.25,y+h*.86,w-.5,h*.12,'浮力方向：竖直向上',15,muted,minimum=12)
    elif data.diagram=='triangle':
        shape(slide,x+w*.22,y+h*.12,w*.57,h*.65,soft,MSO_SHAPE.RIGHT_TRIANGLE,accent)
        box(slide,x+w*.08,y+h*.42,w*.12,.44,'a',19,ink)
        box(slide,x+w*.72,y+h*.34,w*.15,.44,'c',19,ink)
        box(slide,x+.2,y+h*.84,w-.4,h*.14,'a² + b² = c²',20,ink,minimum=14)


def teaching_table(slide,data,theme,x,y,w,h):
    _,ink,accent,surface,soft,_=THEMES[theme]
    if data.title:
        titleh=max(.55,text_lines(data.title,w,16)*16*1.4/72+.06)
        box(slide,x,y,w,titleh,data.title,20,ink,True,minimum=16)
        y+=titleh+.03;h-=titleh+.03
    cells=[data.headers,*data.rows];cols=len(data.headers)
    # Allocate row height from its longest wrapped cell before creating the table.
    # Reject overfull evidence instead of clipping cells or dropping rows.
    weights=[max(sum(1.05 if unicodedata.east_asian_width(c) in 'WF' else .62 for c in row[j]) for row in cells) for j in range(cols)]
    minimum=min(1.5,w/cols)
    widths=[minimum+(w-minimum*cols)*weight/sum(weights) for weight in weights]
    for size in range(20,16,-1):
        heights=[max(.48,max(text_lines(c,widths[j]-.24,size) for j,c in enumerate(row))*size*1.4/72+.12) for row in cells]
        if sum(heights)<=h:break
    else:raise RenderError('text_overflow')
    extra=(h-sum(heights))/len(cells)
    table=slide.shapes.add_table(len(cells),cols,Inches(x),Inches(y),Inches(w),Inches(h)).table
    for j,width in enumerate(widths):table.columns[j].width=Inches(width)
    table.first_row=True;table.horz_banding=False
    for i,row in enumerate(cells):
        table.rows[i].height=Inches(heights[i]+extra)
        for j,text in enumerate(row):
            cell=table.cell(i,j);cell.text=keep_number_unit(text)
            cell.margin_left=cell.margin_right=Inches(.12)
            cell.margin_top=cell.margin_bottom=Inches(.06)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=rgb(('FFE99B' if theme=='classroom' else accent) if i==0 else surface if i%2 else soft)
            cell.text_frame.word_wrap=True;cell.text_frame.auto_size=MSO_AUTO_SIZE.NONE
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(size)
                p.font.bold=i==0;p.font.color.rgb=rgb(THEMES[theme][0] if i==0 and theme!='classroom' else ink)
                p.font._element.append(parse_xml(f'<a:ea {nsdecls("a")} typeface="{FONT}"/>'))
                p.space_after=p.space_before=Pt(0);p.line_spacing=1.15


def teaching_figure(slide,data,theme,x,y,w,h):
    _,ink,accent,_,soft,muted=THEMES[theme]
    if data.title:
        box(slide,x,y,w,.55,data.title,20,ink,True,minimum=16)
        y+=.65;h-=.65
    if data.kind=='interval':
        # Native shapes retain numeric proportions; label lanes with leaders keep
        # close points readable without moving ticks or falsifying the scale.
        count=len(data.items);lane=w/count
        labels=[item.label+('\n'+item.detail if item.detail else '') for item in data.items]
        labelh=max(text_lines(t,lane-.12,17)*17*1.4/72+.05 for t in labels)
        valueh=max(text_lines(f'{b-a:g} {data.unit}',w/(count-1)-.16,18)*18*1.4/72+.05 for a,b in zip(data.values,data.values[1:]))
        if labelh+valueh+1.06>h:raise RenderError('text_overflow')
        axisy=y+labelh+.3
        total=data.values[-1]-data.values[0]
        start=x+.18;length=w-.36
        points=[start+(v-data.values[0])/total*length for v in data.values]
        def leader(x1,y1,x2,y2):
            line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
            line.line.color.rgb=rgb(muted);line.line.width=Pt(.8)
            line._element.spPr.append(parse_xml(f'<a:effectLst {nsdecls("a")}/>'))
        def label_box(tx,ty,tw,th,text,size,bold=True):
            label=box(slide,tx,ty,tw,th,text,size,ink,bold)
            label.text_frame.margin_left=label.text_frame.margin_right=Pt(4)
            for p in label.text_frame.paragraphs:p.alignment=PP_ALIGN.CENTER
        for i,(label,px) in enumerate(zip(labels,points)):
            tx=x+i*lane+.06
            label_box(tx,y,lane-.12,labelh,label,17)
            leader(tx+(lane-.12)/2,y+labelh,px,axisy)
        for i,(a,b) in enumerate(zip(points,points[1:])):
            segment=shape(slide,a,axisy,b-a,.13,accent if i%2==0 else muted)
            segment.name=f'interval-segment-{i+1}'
            tx=x+i*w/(count-1)+.08;tw=w/(count-1)-.16
            leader((a+b)/2,axisy+.13,tx+tw/2,axisy+.38)
            label_box(tx,axisy+.4,tw,valueh,f'{data.values[i+1]-data.values[i]:g} {data.unit}',18)
        for px in points:shape(slide,px-.009,axisy-.06,.018,.25,ink)
        box(slide,x,y+h-.4,w,.4,f'总跨度 {total:g} {data.unit} · 线段按数值比例绘制',15,muted)
        return
    if data.kind=='route':
        box(slide,x,y,w,.32,'路线示意 · 不表示实际距离或方向',11,muted)
        y+=.45;h-=.45
        # A schematic journey, never a fabricated map or a scaled displacement diagram.
        count=len(data.items);cw=w/count
        points=[(x+cw*(i+.5),y+.18+(0 if i%2 else min(.55,h*.25))) for i in range(count)]
        for i in range(count-1):
            start,end=points[i],points[i+1]
            line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(start[0]),Inches(start[1]),Inches(end[0]),Inches(end[1]))
            line.line.color.rgb=rgb(accent);line.line.width=Pt(4)
            line.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle"/>'))
        labely=y+min(.55,h*.25)+.48
        available=y+h-labely
        for i,(item,(px,py)) in enumerate(zip(data.items,points)):
            shape(slide,px-.1,py-.1,.2,.2,THEMES[theme][0],MSO_SHAPE.OVAL,accent)
            tx=x+i*cw+.04;tw=cw-.08
            labelh=max(.42,text_lines(item.label,tw,19)*19*1.4/72+.06)
            box(slide,tx,labely,tw,labelh,item.label,21,ink,True,minimum=19)
            if item.detail:
                box(slide,tx,labely+labelh+.05,tw,available-labelh-.05,item.detail,18,muted,minimum=16)
        return
    if data.kind=='timeline':
        # Positions describe order only; labels explicitly preserve that distinction.
        box(slide,x,y,w,.32,'顺序示意 · 不按时间比例',11,muted)
        y+=.45;h-=.45
    horizontal=w>=8
    count=len(data.items);gap=.3 if horizontal else .15
    cw=(w-gap*(count-1))/count if horizontal else w
    ch=h if horizontal else (h-gap*(count-1))/count
    for i,item in enumerate(data.items):
        px=x+i*(cw+gap) if horizontal else x
        py=y if horizontal else y+i*(ch+gap)
        if data.kind=='timeline':
            shape(slide,px+.04,py+.13,.12,.12,accent,MSO_SHAPE.OVAL)
        else:
            box(slide,px,py,.4,.36,str(i+1),16,accent,True)
        if i<count-1:
            if horizontal:
                start,end=(px+.5 if data.kind=='flow' else px+.24,py+.19),(px+cw+gap-.03,py+.19)
            else:
                start,end=(px+.1,py+.43),(px+.1,py+ch+gap-.04)
            connector=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(start[0]),Inches(start[1]),Inches(end[0]),Inches(end[1]))
            connector.line.color.rgb=rgb(accent);connector.line.width=Pt(2)
        tx=px if horizontal else px+.5
        ty=py+.52 if horizontal else py
        tw=cw if horizontal else cw-.5
        available=ch-.52 if horizontal else ch
        if not horizontal and item.detail and ch<1.15:
            box(slide,tx,ty,tw*.43-.1,available,item.label,20,ink,True,minimum=17)
            box(slide,tx+tw*.43,ty,tw*.57,available,item.detail,19,muted,minimum=17)
            continue
        labelh=max(.4,text_lines(item.label,tw,20)*20*1.4/72+.05) if item.detail else available
        if labelh>available:raise RenderError('text_overflow')
        box(slide,tx,ty,tw,labelh,item.label,22,ink,True,minimum=20)
        if item.detail:
            # Mixed algebra/CJK and emphasized runs wrap wider in Office than
            # the estimate. Keep a readable detail size in narrow flow columns.
            box(slide,tx,ty+labelh+.06,tw,available-labelh-.06,item.detail,17 if horizontal and tw<3 else 20,muted,minimum=17)


def teaching_picture(slide,data,theme,x,y,w,h):
    from PIL import Image
    _,ink,_,_,_,_=THEMES[theme]
    raw=base64.b64decode(data.data_url.split(',',1)[1],validate=True)
    with Image.open(io.BytesIO(raw)) as im:pw,ph=im.size
    captionh=text_lines(data.caption,w,16)*16*1.4/72+.1
    area=h-captionh-.12
    if area<.8:raise RenderError('text_overflow')
    scale=min(w/pw,area/ph);iw,ih=pw*scale,ph*scale
    picture=slide.shapes.add_picture(io.BytesIO(raw),Inches(x+(w-iw)/2),Inches(y+(area-ih)/2),width=Inches(iw),height=Inches(ih))
    picture._element.nvPicPr.cNvPr.set('descr',data.caption)
    box(slide,x,y+area+.12,w,captionh,data.caption,18,ink,minimum=16)


def emphasize(slide,terms,theme):
    if not terms:return
    terms=[keep_number_unit(t) for t in terms]
    pattern=re.compile('('+ '|'.join(re.escape(t) for t in sorted(set(terms),key=len,reverse=True))+')')
    color='FFD166' if theme=='midnight' else 'B54822'
    for sh in slide.shapes:
        frames=[(sh.text_frame,False)] if sh.has_text_frame else []
        if sh.has_table:frames.extend((c.text_frame,i==0) for i,row in enumerate(sh.table.rows) for c in row.cells)
        for frame,header in frames:
            for p in frame.paragraphs:
                text=p.text
                if not pattern.search(text):continue
                p.clear()
                for part in pattern.split(text):
                    if not part:continue
                    run=p.add_run();run.text=part
                    if part in terms:
                        run.font.bold=True
                        if header:run.font.underline=True
                        else:run.font.color.rgb=rgb(color)


def formula(slide,text,theme,editable=False,x=.7,y=6.13,width=11.9,band=True):
    # Figure + Agg canvas avoids pyplot's process-global current-figure state.
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib import rc_context
    from PIL import Image
    _,ink,accent,_,soft,_=THEMES[theme]
    fig=Figure(figsize=(5,.65));FigureCanvasAgg(fig)
    # Model JSON may escape an already escaped command; leave non-command
    # backslashes alone (e.g. LaTeX row breaks and escaped punctuation).
    text=re.sub(r'\\{2,}(?=[A-Za-z])',lambda match:'\\',text)
    aliases={'implies':'Rightarrow','impliedby':'Leftarrow','iff':'Leftrightarrow'}
    text=re.sub(r'\\([A-Za-z]+)',lambda m:'\\'+aliases.get(m[1],m[1]),text)
    pic=io.BytesIO()
    face=mathtext_font()
    with rc_context({'font.family':face,'mathtext.fontset':'custom',
                     'mathtext.rm':face,'mathtext.it':face,'mathtext.bf':face}):
        fig.text(.02,.3,'$'+text.strip('$')+'$',fontsize=24,color='#'+ink)
        fig.savefig(pic,format='png',transparent=True,bbox_inches='tight',dpi=160,pad_inches=.03)
    pic.seek(0)
    with Image.open(pic) as im:pw,ph=im.size
    # Keep the native 24pt formula scale unless a modest reduction is enough.
    scale=min(1,.53/(ph/160),(width-1)/(pw/160))
    if scale<.68:raise RenderError('formula_overflow')
    w=pw/160*scale;h=ph/160*scale
    band_height=.8 if editable else .65
    if band:
        shape(slide,x,y,width,band_height,soft,MSO_SHAPE.ROUNDED_RECTANGLE)
        shape(slide,x,y,.045,band_height,accent)
    if editable:
        from latex2mathml.converter import convert
        import mathml2omml
        equation=mathml2omml.convert(convert(text.strip('$')))
        namespace='http://schemas.openxmlformats.org/officeDocument/2006/math'
        drawing='http://schemas.microsoft.com/office/drawing/2010/main'
        word='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        element=parse_xml(f'<a14:m xmlns:a14="{drawing}" xmlns:m="{namespace}" xmlns:w="{word}"><m:oMathPara>{equation}</m:oMathPara></a14:m>')
        # Office uses Word run properties inside its editable equation objects.
        for run in element.findall('.//{'+namespace+'}r'):
            style=parse_xml(f'<w:rPr xmlns:w="{word}"><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math" w:eastAsia="Microsoft YaHei"/><w:sz w:val="{round(48*scale)}"/><w:color w:val="{ink}"/></w:rPr>')
            run.insert(1,style)
        # Native Office fractions occupy more height than the raster preview;
        # retain the math font and reserve room through the footer's safe gap.
        # Native Office math is wider than its raster estimate. The old 1-inch
        # reserve clipped correct sidebar equations despite space in the band.
        frame=slide.shapes.add_textbox(Inches(x+.28),Inches(y+.01),Inches(width-.42),Inches(.78)).text_frame
        frame.margin_left=frame.margin_top=frame.margin_bottom=0
        frame.paragraphs[0]._p.append(element)
    else:
        pic.seek(0);slide.shapes.add_picture(pic,Inches(x+.28),Inches(y+.06),width=Inches(w),height=Inches(h))


def render(deck,start_index=0,total_pages=None):
    total_pages=total_pages or len(deck.slides)
    ppt=Presentation();ppt.slide_width=Inches(13.333);ppt.slide_height=Inches(7.5)
    bg,ink,accent,_,soft,muted=THEMES[deck.theme]
    for index,data in enumerate(deck.slides,start_index):
        field='bullets'
        try:
            if any(len(t)>160 for t in data.bullets):raise RenderError('text_overflow')
            layout=choose_layout(data,index,total_pages)
            s=ppt.slides.add_slide(ppt.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(bg)
            teaching_media=bool(data.table or data.figure)
            paired_picture=bool(data.picture and teaching_media)
            has_visual=bool(data.scientific or data.graph or data.chart or data.diagram or teaching_media or data.picture)
            cover=layout=='cover' and not(data.graph or data.scientific)
            # Teaching content supplies the visual evidence; a cover without a
            # real diagram must not acquire a fake graph or large empty panels.
            decorate(s,deck.theme,False)
            # Full deck title remains in document metadata; a compact running header
            # must not crowd the slide title when the deck name is 100 CJK characters.
            running=deck.title if len(deck.title)<=32 else deck.title[:31]+'…'
            classroom=deck.theme=='classroom'
            field='section'
            if classroom:
                box(s,.8,.56,10.9,.65,data.section or LAYOUT_LABELS[layout],26,accent,True,minimum=15)
            else:
                box(s,1.55 if deck.theme=='coral' else .7,.37,7,.34,'备课工坊 / '+running,11,muted,minimum=10)
                box(s,10.1,.95,2.5,.3,LAYOUT_LABELS[layout],12,accent,True)
                if data.section:box(s,.7,.94,9.1,.43,data.section,12,muted,minimum=10)
            titlefont=('SimSun' if os.name=='nt' else 'Noto Serif CJK SC') if deck.theme=='paper' else FONT
            field='title'
            specialist=bool(data.graph or data.scientific)
            evidence_first=specialist and layout in ('concept','comparison','summary')
            focusx=8.2 if specialist and evidence_first else .7
            focusw=4.4 if specialist and evidence_first else 5.45 if specialist else 11.9
            # A long question must wrap within its text column while the
            # evidence keeps its area. Reserve actual readable line height.
            focush=max(.75,text_lines(data.focus,focusw-.23,18)*18*1.4/72+.07) if specialist else .55
            bottom=6.0 if data.formula else 6.78
            if data.focus:bottom-=focush-.03 if specialist else .72
            # Several scientific relations separated by TeX spacing need a
            # shared equation row. Keep the complete expressions at readable
            # scale rather than squeezing them into the text column.
            wide_science_formula=bool(data.scientific and r'\quad' in data.formula)
            if teaching_media:
                object_region(s,'title',box,.7,1.43,11.9,.89,data.title,32,ink,True,minimum=22,font=titlefont)
                field='bullets'
                caption='\n'.join(data.bullets)
                dense_table=bool(data.table and len(data.table.rows)>=4 and data.formula and not(data.figure or data.picture))
                if dense_table:
                    # Put prose beside complete experimental records, preserving units.
                    object_region(s,'body',box,.7,2.4,3.5,bottom-2.4,caption,23,ink,minimum=19)
                    field='table'
                    object_region(s,'visual',teaching_table,data.table,deck.theme,4.5,2.4,8.1,bottom-2.4)
                else:
                    captionh=text_lines(caption,11.9,19)*19*1.4/72+.08 if caption else 0
                    mediah=bottom-2.5-captionh-(.18 if caption else 0)
                    if mediah<2:raise RenderError('text_overflow')
                    field='table' if data.table else 'figure'
                    # Uploaded illustrations are supporting material; give a comparison
                    # table enough width for readable CJK cells before shrinking fonts.
                    tablew=8.25 if data.picture else 6.25 if data.figure else 11.9
                    if data.table:
                        object_region(s,'visual',teaching_table,data.table,deck.theme,.7,2.5,tablew,mediah)
                    if data.figure:
                        field='figure'
                        object_region(s,'visual',teaching_figure,data.figure,deck.theme,7.35 if data.table else .7,2.5,5.25 if data.table else 6.25 if data.picture else 11.9,mediah)
                    if paired_picture:
                        field='picture'
                        object_region(s,'visual',teaching_picture,data.picture,deck.theme,9.35 if data.table else 7.35,2.5,3.25 if data.table else 5.25,mediah)
                    if caption:
                        field='bullets'
                        object_region(s,'body',box,.7,bottom-captionh,11.9,captionh,caption,21,ink,minimum=19)
            elif cover:
                object_region(s,'title',box,.7,1.65,7.1 if has_visual else 11.9,1.82,data.title,40,ink,True,minimum=22,font=titlefont)
                field='bullets'
                object_region(s,'body',body,data,layout,deck.theme,.8,3.75,6.65 if has_visual else 10.9,bottom-3.85)
            else:
                object_region(s,'title',box,.7,1.43,11.9,.89,data.title,30,ink,True,minimum=21,font=titlefont)
                field='bullets'
                # Evidence comes first for explanation/comparison; procedures
                # retain a sequence and exercises retain their question panel.
                copy_layout=('comparison' if specialist and layout=='comparison' else 'concept' if evidence_first else 'steps' if specialist and layout in ('example','cover') else layout)
                object_region(s,'body',body,data,copy_layout,deck.theme,8.2 if evidence_first else .7,2.55,4.4 if evidence_first else 5.45 if specialist else 6.9 if has_visual else 11.9,bottom-2.66)
            if data.picture and not teaching_media:
                field='picture'
                object_region(s,'visual',teaching_picture,data.picture,deck.theme,7.95,1.65 if cover else 2.55,4.65,bottom-(1.76 if cover else 2.66))
            elif has_visual and not teaching_media:
                field='scientific' if data.scientific else 'graph' if data.graph else 'chart' if data.chart else 'diagram'
                object_region(s,'visual',visual,data,deck.theme,.7 if evidence_first else 6.5 if specialist else 7.95,2.55,7.15 if evidence_first else 6.1 if specialist else 4.65,3.45 if wide_science_formula else 4.12 if specialist else bottom-2.66)
            if data.focus:
                field='focus'
                focusy=bottom+.02 if specialist else bottom+.12
                if classroom:object_region(s,'focus',shape,focusx,focusy-.02,focusw,focush+.07,'FFF2B3',MSO_SHAPE.ROUNDED_RECTANGLE)
                object_region(s,'focus',shape,focusx,focusy,.05,focush,accent)
                object_region(s,'focus',box,focusx+.23,focusy,focusw-.23,focush,data.focus,23,ink,True,minimum=18)
            if data.formula:
                field='formula'
                object_region(s,'formula',formula,data.formula,deck.theme,deck.editable_formulas,x=8.2 if specialist and evidence_first and not wide_science_formula else .7,width=11.9 if wide_science_formula else 4.4 if specialist and evidence_first else 5.45 if specialist else 11.9,band=layout in ('concept','summary') and not specialist)
            for key,items in region_shapes(s).items():
                for item in items:
                    node=item._element.xpath('.//p:cNvPr')[0]
                    allowed=key!='formula' and not(key=='visual' and (data.graph or data.scientific or data.chart or data.diagram))
                    node.set('descr',node.get('descr','')+'|edu-rotatable='+str(int(allowed)))
            apply_canvas_positions(s,data.canvas_positions,ppt.slide_width,ppt.slide_height,data)
            for obj in data.canvas_objects:draw_canvas_object(s,obj,ppt.slide_width,ppt.slide_height)
            emphasize(s,data.emphasis,deck.theme)
            shape(s,.7,7.4,11.9*(index+1)/total_pages,.035,accent)
            box(s,.7,7.06,9,.25,'待核验草稿 · 请教师核对' if deck.needs_review else '教学资料 · 来源详见备注',10,muted)
            box(s,11.5,7.04,1.1,.3,f'{index+1:02} / {total_pages:02}',11,accent)
            s.notes_slide.notes_text_frame.text=data.notes
            if data.answer:s.notes_slide.notes_text_frame.text+='\n课堂答案（演示时由教师揭示）：\n'+data.answer
            if data.picture:
                s.notes_slide.notes_text_frame.text+='\n配图说明：'+data.picture.caption+'\n图片来源（教师填写）：'+data.picture.source+'\n图片内容与使用权限请教师核对；文字评审未检查图片。'
            if data.canvas_order:apply_canvas_order(s,data.canvas_order)
        except ValueError as e:
            code=e.code if isinstance(e,RenderError) else ('formula_invalid' if field=='formula' else 'content_invalid')
            raise RenderError(code,field,index+1) from e
    ppt.core_properties.title=deck.title
    result=io.BytesIO();ppt.save(result);return result.getvalue()


@app.get('/health')
def health():return {'status':'ok','courseware_canvas_version':'canvas-v3','pptx':True,'generated_code_execution':False,'courseware_media_version':'media-v1','courseware_layout_version':'page-fit-v1','courseware_composition_version':'teaching-purpose-v2','courseware_plot_version':'matplotlib-v1','courseware_graph_version':'graph-v1','scientific_diagrams_version':'science-v1','pdf_extraction_version':'math-layout-v2'}

@app.post('/render/object-map')
async def object_map(request:Request,page:int=1):
    raw=bytearray()
    async for part in request.stream():
        raw.extend(part)
        if len(raw)>20*1024*1024:raise HTTPException(413,'PPTX too large')
    try:
        # Stored render output only; impose archive expansion limits before parsing.
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist())>5000 or sum(i.file_size for i in archive.infolist())>100*1024*1024:raise ValueError('archive capacity')
        ppt=Presentation(io.BytesIO(raw))
        if not 1<=page<=len(ppt.slides):raise ValueError('page')
        objects=[]
        for key,items in region_shapes(ppt.slides[page-1]).items():
            x=min(i.left for i in items);y=min(i.top for i in items)
            right=max(i.left+i.width for i in items);bottom=max(i.top+i.height for i in items)
            fields=dict(part.split('=',1) for part in items[0]._element.xpath('.//p:cNvPr')[0].get('descr','').split('|') if '=' in part)
            if 'edu-box' in fields:
                x,y,w,h=[int(v) for v in fields['edu-box'].split(',')];right=x+w;bottom=y+h
            minimum=float(fields['edu-base-min'])/float(fields['edu-scale']) if 'edu-base-min' in fields and 'edu-scale' in fields else canvas_min_scale(items,key)
            objects.append({'id':key,'x':x/ppt.slide_width,'y':y/ppt.slide_height,'width':(right-x)/ppt.slide_width,'height':(bottom-y)/ppt.slide_height,'min_scale':minimum,'rotatable':fields.get('edu-rotatable')=='1','rotation':float(fields.get('edu-rotation',0))})
        return {'objects':objects,'version':'canvas-v3'}
    except Exception:raise HTTPException(422,'Invalid stored PPTX or page')


@app.post('/render/pptx')
def pptx(deck:Deck):
    try:return Response(render(deck),media_type='application/vnd.openxmlformats-officedocument.presentationml.presentation')
    except RenderError as e:raise HTTPException(422,{'code':e.code,'page':e.page,'field':e.field})

@app.post('/render/canvas-preview-pptx')
def canvas_preview_pptx(deck:Deck,page:int=1):
    if not 1<=page<=len(deck.slides):raise HTTPException(422,'Invalid page')
    total=len(deck.slides)
    one=deck.model_copy(update={'slides':[deck.slides[page-1]]})
    try:return Response(render(one,page-1,total),media_type='application/vnd.openxmlformats-officedocument.presentationml.presentation')
    except RenderError as e:raise HTTPException(422,{'code':e.code,'page':e.page,'field':e.field})


@app.post('/extract')
async def extract(file:UploadFile):
    data=await file.read(100*1024*1024+1)
    if len(data)>100*1024*1024:raise HTTPException(413,'文件超过100MB')
    suffix=Path(file.filename or '').suffix.lower()
    try:
        if suffix in ('.docx','.pptx'):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if len(z.infolist())>5000 or sum(x.file_size for x in z.infolist())>100*1024*1024:raise ValueError('文档解压后过大')
        if suffix in ('.txt','.md'):
            pages=[{'page':1,'text':data.decode('utf-8-sig')}]
        elif suffix=='.pdf':
            from pdf_extract import extract_pdf
            pages=extract_pdf(data)
        elif suffix=='.docx':
            from docx import Document
            d=Document(io.BytesIO(data));texts=[p.text for p in d.paragraphs]
            texts += [' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows]
            pages=[{'page':1,'text':'\n'.join(texts)}]
        elif suffix=='.pptx':
            p=Presentation(io.BytesIO(data));pages=[{'page':i+1,'text':'\n'.join(s.text for s in sl.shapes if s.has_text_frame)} for i,sl in enumerate(p.slides)]
        else:raise ValueError('支持PDF/DOCX/PPTX/TXT/MD')
        # Some PDF glyph separators extract as NUL, which PostgreSQL cannot store.
        for page in pages:page['text']=page['text'].replace('\x00',' ')
        if not any(p['text'].strip() for p in pages):raise ValueError('未提取到文字；首期不支持扫描件OCR')
        if sum(len(p['text']) for p in pages)>1500000:raise ValueError('提取文字超过上限')
        return {'pages':pages}
    except Exception as e:raise HTTPException(422,'文档解析失败：'+str(e)[:200])

@app.post('/reference/extract')
async def reference_extract(file:UploadFile):
    from reference_extract import inspect_reference
    data=await file.read(20*1024*1024+1)
    if len(data)>20*1024*1024:raise HTTPException(413,'参考课件超过20MB')
    try:return inspect_reference(data,Path(file.filename or '').suffix.lower())
    except Exception:raise HTTPException(422,'参考课件无法读取画面；请使用未加密的PDF/PPTX，最多80页，并检查画面解析服务')


@app.post('/document/page-image')
async def document_page_image(file:UploadFile,page:int=Query(ge=1,le=300)):
    data=await file.read(100*1024*1024+1)
    if len(data)>100*1024*1024:raise HTTPException(413,'文件超过100MB')
    try:
        import pymupdf
        if Path(file.filename or '').suffix.lower()!='.pdf':raise ValueError()
        with pymupdf.open(stream=data,filetype='pdf') as doc:
            if doc.is_encrypted or len(doc)>300 or page>len(doc):raise ValueError()
            original=doc[page-1]
            scale=1600/max(original.rect.width,original.rect.height)
            png=original.get_pixmap(matrix=pymupdf.Matrix(scale,scale),alpha=False).tobytes('png')
            return {'page':page,'data_url':'data:image/png;base64,'+base64.b64encode(png).decode()}
    except Exception:raise HTTPException(422,'原始PDF页面无法读取，请检查页码或文件')

class FormulaRequest(BaseModel):
    formula:str=Field(min_length=1,max_length=250)

@app.post('/render/formula')
def formula_preview(request:FormulaRequest):
    try:
        p=Presentation();sl=p.slides.add_slide(p.slide_layouts[6]);formula(sl,request.formula,'ocean')
        picture=next(s for s in sl.shapes if s.shape_type==13)
        return Response(picture.image.blob,media_type='image/png')
    except Exception:raise HTTPException(422,{'code':'formula_invalid','field':'formula'})
