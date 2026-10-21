import io

from fastapi.testclient import TestClient
from pptx import Presentation

from app import app


def test_function_series_export_as_editable_click_groups():
    graph = dict(title='比较开口方向', x_min=-3, x_max=3, y_min=-4, y_max=4,
                 reveal_on_click=True, series=[dict(kind='quadratic', a=1, mark_x=[1]),
                                               dict(kind='quadratic', a=-1, mark_x=[-1])])
    response = TestClient(app).post('/render/pptx', json=dict(title='原生点击动画', slides=[
        dict(title='每次点击出现一条曲线', graph=graph)]))
    assert response.status_code == 200, response.text
    deck = Presentation(io.BytesIO(response.content))
    slide = deck.slides[0]
    groups = [s for s in slide.shapes if s.name.startswith('function-series-')]
    assert len(groups) == 2, 'each curve, points and legend must form an editable group'
    assert all(len(s.shapes) > 10 for s in groups)
    assert not any(s.shape_type == 13 for s in slide.shapes)
    assert len(slide._element.xpath('./p:timing')) == 1
    effects = slide._element.xpath('.//p:cTn[@nodeType="clickEffect"]')
    assert len(effects) == 2
    targets = slide._element.xpath('.//p:timing//p:spTgt/@spid')
    assert targets == [str(s.shape_id) for s in groups]
    assert slide._element.xpath('.//p:timing//p:strVal/@val') == ['visible', 'visible']
    assert len(slide._element.xpath('.//p:timing//p:cond[@evt="onNext"]')) == 1
    from test_render import assert_inside_slide
    assert_inside_slide(deck)


def test_click_animation_rejects_raster_or_gif_combination():
    graph = dict(title='互斥模式', x_min=-3, x_max=3, y_min=-4, y_max=4,
                 reveal_on_click=True, series=[dict(kind='linear', k=1)])
    client = TestClient(app)
    for extra in [dict(renderer='matplotlib'), dict(animation=dict(series=0, param='k', to=2, duration_ms=2000))]:
        response = client.post('/render/pptx', json=dict(title='校验', slides=[dict(title='模式', graph={**graph, **extra})]))
        assert response.status_code == 422
