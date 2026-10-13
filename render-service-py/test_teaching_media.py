import io
import base64
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from app import app
from test_render import assert_inside_slide, THEMES
import json
from pathlib import Path

TABLE = {'title': '列车时刻表（教学示例）', 'headers': ['站点', '到站', '离站'],
         'rows': [['甲站', '—', '08:00'], ['乙站', '08:30', '08:32'], ['丙站', '09:00', '—']]}

def test_five_experiment_records_fit_beside_caption_and_formula():
    table={'title':'实验记录（示意数据，非实测，以课堂实测为准）',
           'headers':['浸入情况','F拉 / N','F浮 / N'],
           'rows':[['未浸入','2.7','0'],['浸入约1/4体积（水）','2.45','0.25'],
                   ['浸入约1/2体积（水）','2.20','0.50'],['完全浸没（水中）','1.70','1.00'],['完全浸没（盐水）','1.50','1.20']]}
    caption='同一铝块，固定排液体积后比较不同液体，记录并计算浮力。'
    r=TestClient(app).post('/render/pptx',json={'title':'浮力','theme':'ocean','editable_formulas':True,
        'slides':[{'title':'探究浮力的影响因素','table':table,'formula':'F_{浮}=G-F_{拉}','bullets':[caption]}]})
    assert r.status_code==200,r.text
    p=Presentation(io.BytesIO(r.content));assert_inside_slide(p)
    actual=next(s.table for s in p.slides[0].shapes if s.has_table)
    assert [[c.text.replace('\u00a0',' ') for c in row.cells] for row in actual.rows]==[table['headers'],*table['rows']]
    assert any(s.has_text_frame and s.text==caption for s in p.slides[0].shapes)
    assert '示意数据' in '\n'.join(s.text for s in p.slides[0].shapes if s.has_text_frame)
    assert '<m:oMath' in p.slides[0]._element.xml

@pytest.mark.parametrize('layout', ['comparison', 'exercise'])
def test_short_teaching_content_has_no_full_height_color_panel(layout):
    r = TestClient(app).post('/render/pptx', json={'title': '观察与推理', 'theme': 'ocean', 'slides': [
        {'title': '解释变化规律', 'layout': layout, 'bullets': ['x增加1，y增加2。', '斜率表示变化率。']}]})
    assert r.status_code == 200, r.text
    ppt = Presentation(io.BytesIO(r.content))
    panels = [s for s in ppt.slides[0].shapes if s.shape_type == 1 and s.width > 3*914400 and s.height > 1.8*914400 and s.top > 1*914400]
    assert not panels, 'Short content should use readable open text, not an empty colored panel'

def test_native_formula_is_editable_and_preserves_fraction_and_subscript():
    r = TestClient(app).post('/render/pptx', json={'title': '阿基米德原理', 'editable_formulas': True, 'slides': [
        {'title': '排开液体的体积', 'formula': r'F_{浮}=\rho_{液}gV_{排}=\frac{G}{2}', 'bullets': ['明确液体密度与排液体积。']}]})
    assert r.status_code == 200, r.text
    ppt = Presentation(io.BytesIO(r.content))
    xml = ppt.slides[0]._element.xml
    assert '<m:oMath' in xml and '<m:f>' in xml and '<m:sSub>' in xml
    assert '浮' in xml and '排' in xml
    assert not any(s.shape_type == 13 for s in ppt.slides[0].shapes)

def test_caption_crowding_reports_bullets_instead_of_title():
    response = TestClient(app).post('/render/pptx', json={
        'title': '一次函数', 'theme': 'classroom', 'slides': [{
            'title': '一次函数的图像', 'layout': 'concept', 'table': TABLE,
            'focus': '图像是一条直线', 'formula': 'y=2x+1',
            'bullets': ['观察函数图像并结合数据解释变化规律。'*4]*3}]})
    assert response.status_code == 422
    assert response.json()['detail'] == {'code': 'text_overflow', 'page': 1, 'field': 'bullets'}

def test_long_section_reports_section_instead_of_bullets():
    response = TestClient(app).post('/render/pptx', json={
        'title': '一次函数', 'theme': 'classroom', 'slides': [{
            'title': '图像', 'section': '探究\n证据\n归纳\n应用', 'bullets': ['观察函数图像']}]})
    assert response.status_code == 422
    assert response.json()['detail']['field'] == 'section'

def test_continuation_pages_export_and_remain_bounded():
    from app import Deck
    data = {'title': '自动拆页课件', 'slides': [
        {'title': f'观察与解释 {i+1}', 'bullets': ['观察现象并说明依据']} for i in range(21)]}
    response = TestClient(app).post('/render/pptx', json=data)
    assert response.status_code == 200, response.text
    ppt = Presentation(io.BytesIO(response.content))
    assert len(ppt.slides) == 21
    assert_inside_slide(ppt)
    data['slides'] = data['slides'][:1] * 100
    assert len(Deck.model_validate(data).slides) == 100
    data['slides'].append(data['slides'][0])
    assert TestClient(app).post('/render/pptx', json=data).status_code == 422

def test_interval_uses_numeric_spacing_and_editable_segments():
    interval={'kind':'interval','title':'从出发开始计时','unit':'min','values':[0,30,32,60],
              'items':[{'label':'08:00出发'},{'label':'08:30到站'},{'label':'08:32离站'},{'label':'09:00到达'}]}
    r=TestClient(app).post('/render/pptx',json={'title':'时间间隔','theme':'classroom','slides':[
        {'title':'全程包含哪些过程？','figure':interval,'bullets':['停站2 min，运行58 min。']}]})
    assert r.status_code==200,r.text
    p=Presentation(io.BytesIO(r.content));assert_inside_slide(p)
    segments=[s for s in p.slides[0].shapes if s.name.startswith('interval-segment-')]
    assert len(segments)==3
    assert abs(segments[0].width/segments[1].width-15)<.001
    assert abs(segments[2].width/segments[1].width-14)<.001
    text='\n'.join(s.text for s in p.slides[0].shapes if s.has_text_frame).replace('\u00a0',' ')
    for required in ['30 min','2 min','28 min','60 min','08:30到站']:
        assert required in text
    assert not any(s.shape_type==13 for s in p.slides[0].shapes)

@pytest.mark.parametrize('values',[[0,0],[2,1],[0,1,2],[0,float('inf')]])
def test_interval_rejects_invalid_numeric_positions(values):
    from app import TeachingFigure
    with pytest.raises(ValueError):
        TeachingFigure(kind='interval',unit='min',values=values,items=[{'label':'开始'},{'label':'结束'}])

def test_number_unit_does_not_break_across_lines():
    r=TestClient(app).post('/render/pptx',json={'title':'排版','slides':[{'title':'计算结果',
        'bullets':['全程60 min，速度2.5 m/s，距离100 km。'],'emphasis':['60 min']}]})
    assert r.status_code==200
    p=Presentation(io.BytesIO(r.content))
    body=next(s for s in p.slides[0].shapes if s.has_text_frame and '全程' in s.text)
    assert '60\u00a0min' in body.text and '2.5\u00a0m/s' in body.text and '100\u00a0km' in body.text
    assert any(r.text=='60\u00a0min' and r.font.bold for para in body.text_frame.paragraphs for r in para.runs)
FIGURE = {'kind': 'timeline', 'title': '行程节点（不按比例）', 'items': [
    {'label': '08:00', 'detail': '甲站出发'}, {'label': '08:30', 'detail': '乙站到达'},
    {'label': '09:00', 'detail': '丙站到达'}]}

def test_classroom_route_is_editable_and_headers_keep_contrast():
    route={'kind':'route','title':'一次出行的三个位置','items':[
        {'label':'甲站','detail':'起点'}, {'label':'乙站','detail':'中途停留'},
        {'label':'丙站','detail':'终点'}]}
    r=TestClient(app).post('/render/pptx',json={'title':'运动的描述','theme':'classroom','slides':[
        {'title':'一次出行怎样记录？','table':TABLE,'figure':route,'emphasis':['到站','08:30'],
         'focus':'哪个记录对应瞬间，哪个记录对应过程？'}]})
    assert r.status_code==200,r.text
    p=Presentation(io.BytesIO(r.content));assert_inside_slide(p)
    slide=p.slides[0]
    text='\n'.join(s.text for s in slide.shapes if s.has_text_frame)
    assert '路线示意 · 不表示实际距离或方向' in text
    assert all(item['label'] in text and item['detail'] in text for item in route['items'])
    assert not any(s.shape_type==13 for s in slide.shapes)
    table=next(s.table for s in slide.shapes if s.has_table)
    header=table.cell(0,1).text_frame.paragraphs[0]
    assert header.runs[0].font.color.type is None, 'Header emphasis must inherit contrasting header ink'
    assert any(run.text=='08:30' and run.font.bold for run in table.cell(2,1).text_frame.paragraphs[0].runs)

def test_picture_round_trip_aspect_ratio_source_and_emphasis():
    from PIL import Image
    buf=io.BytesIO();Image.new('RGB',(400,200),'blue').save(buf,format='PNG')
    picture={'data_url':'data:image/png;base64,'+base64.b64encode(buf.getvalue()).decode(),
             'caption':'图像说明', 'source':'教师自有图片'}
    r=TestClient(app).post('/render/pptx',json={'title':'读图', 'slides':[
        {'title':'比较时刻', 'table':TABLE, 'picture':picture, 'emphasis':['08:30'],
         'bullets':['比较08:30和停留2分钟'], 'focus':'这里的时间表示什么？'}]})
    assert r.status_code==200,r.text
    p=Presentation(io.BytesIO(r.content));assert_inside_slide(p)
    pics=[s for s in p.slides[0].shapes if s.shape_type==13]
    assert len(pics)==1
    assert abs(pics[0].width/pics[0].height-2)<.001
    assert '教师自有图片' in p.slides[0].notes_slide.notes_text_frame.text
    text=next(s for s in p.slides[0].shapes if s.has_text_frame and s.text=='比较08:30和停留2分钟')
    assert any(r.text=='08:30' and r.font.bold for para in text.text_frame.paragraphs for r in para.runs)

def test_concept_table_and_uploaded_picture_preserve_readable_cells():
    from PIL import Image
    deck=json.loads((Path(__file__).resolve().parents[1]/'fixtures/courseware-teaching-media.json').read_text(encoding='utf-8'))
    slide=deck['slides'][2]
    buf=io.BytesIO();Image.new('RGB',(400,200),'white').save(buf,format='PNG')
    slide['picture']={'data_url':'data:image/png;base64,'+base64.b64encode(buf.getvalue()).decode(),
                      'caption':'时刻与时间间隔示意','source':'教师自绘测试素材'}
    r=TestClient(app).post('/render/pptx',json={'title':deck['title'],'theme':'classroom','slides':[slide]})
    assert r.status_code==200,r.text
    p=Presentation(io.BytesIO(r.content));assert_inside_slide(p)
    table=next(s for s in p.slides[0].shapes if s.has_table)
    picture=next(s for s in p.slides[0].shapes if s.shape_type==13)
    assert table.left+table.width<picture.left
    assert all(p.font.size.pt>=17 for row in table.table.rows for c in row.cells for p in c.text_frame.paragraphs)
    assert [[c.text for c in row.cells] for row in table.table.rows]==[slide['table']['headers'],*slide['table']['rows']]

@pytest.mark.parametrize('data', ['https://example.com/private.png','file:///secret',
    'data:image/svg+xml;base64,PHN2Zz4=', 'data:image/png;base64,bm90IGFuIGltYWdl'])
def test_picture_rejects_urls_and_invalid_bytes(data):
    r=TestClient(app).post('/render/pptx',json={'title':'检查','slides':[{'title':'检查',
        'picture':{'data_url':data,'source':'教师自有图片','caption':'示意'}}]})
    assert r.status_code==422
    assert 'data_url' not in r.text and 'secret' not in r.text

@pytest.mark.parametrize('theme', THEMES)
def test_editable_teaching_evidence_and_focus(theme):
    slide = {'title': '时刻与时间间隔', 'layout': 'comparison',
             'bullets': ['观察到站、离站和停留时间的区别'], 'table': TABLE, 'figure': FIGURE,
             'focus': '08:30是时刻，停留2分钟是时间间隔。', 'notes': '数据为教学示例，非真实车次。'}
    r = TestClient(app).post('/render/pptx', json={'title': '时间与位移', 'theme': theme, 'slides': [slide]})
    assert r.status_code == 200, r.text
    p = Presentation(io.BytesIO(r.content)); assert_inside_slide(p)
    tables = [s.table for s in p.slides[0].shapes if s.has_table]
    assert len(tables) == 1
    assert [[c.text for c in row.cells] for row in tables[0].rows] == [TABLE['headers'], *TABLE['rows']]
    text = '\n'.join(s.text for s in p.slides[0].shapes if s.has_text_frame)
    for value in [slide['focus'], FIGURE['title'], *[i['detail'] for i in FIGURE['items']]]:
        assert value in text
    assert p.slides[0].notes_slide.notes_text_frame.text == slide['notes']
    assert not any(s.shape_type == 13 for s in p.slides[0].shapes), 'Evidence must remain editable'

@pytest.mark.parametrize('extra', [
    {'table': {**TABLE, 'rows': [['only one cell']]}},
    {'table': {**TABLE, 'headers': ['', '到站', '离站']}},
    {'table': {**TABLE, 'rows': [['密集文字'*20, '08:00', '08:30']]}},
    {'figure': {**FIGURE, 'kind': 'invented'}},
    {'figure': {**FIGURE, 'items': [{'label': '', 'detail': '内容'}]*3}},
    {'table': TABLE, 'chart': {'labels': ['A'], 'values': [1]}},
    {'figure': FIGURE, 'diagram': 'triangle'},
    {'focus': '过长'*31},
])
def test_media_invalid_is_rejected_without_private_input(extra):
    r = TestClient(app).post('/render/pptx', json={'title': '私密标题', 'slides': [{'title': '检查', **extra}]})
    assert r.status_code == 422
    assert r.json()['detail']['page'] == 1
    assert '私密标题' not in r.text

@pytest.mark.parametrize('theme', THEMES)
def test_reference_structure_fixture_preserves_all_content(theme):
    deck=json.loads((Path(__file__).resolve().parents[1]/'fixtures/courseware-teaching-media.json').read_text(encoding='utf-8'))
    deck['theme']=theme
    r=TestClient(app).post('/render/pptx',json=deck)
    assert r.status_code==200,r.text
    ppt=Presentation(io.BytesIO(r.content));assert_inside_slide(ppt)
    for source,slide in zip(deck['slides'],ppt.slides):
        text='\n'.join(s.text for s in slide.shapes if s.has_text_frame).replace('\u00a0',' ')
        for required in [source['title'],*source['bullets'],source.get('focus','')]:
            assert required in text
        assert source['notes']==slide.notes_slide.notes_text_frame.text
