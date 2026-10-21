"""Bounded reference inspection. No generated code, network, or persistent uploads."""
import base64
import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pdf_extract import extract_pdf

def inspect_reference(data,suffix):
    images=[]
    if suffix=='.pdf':
        pages=extract_pdf(data,max_pages=80)
        selected=sorted(set(round(i*(len(pages)-1)/min(5,len(pages)-1))+1 for i in range(min(6,len(pages))))) if len(pages)>1 else [1]
        executable=shutil.which('pdftoppm')
        if not executable:raise ValueError('PDF image renderer unavailable')
        with tempfile.TemporaryDirectory(prefix='edu-reference-') as folder:
            path=Path(folder)/'reference.pdf';path.write_bytes(data)
            for n in selected:
                prefix=Path(folder)/f'page-{n}'
                subprocess.run([executable,'-f',str(n),'-l',str(n),'-scale-to','1280','-singlefile','-png',str(path),str(prefix)],check=True,capture_output=True,timeout=12)
                with Image.open(str(prefix)+'.png') as im:
                    im=im.convert('RGB');buf=io.BytesIO();im.save(buf,'JPEG',quality=80)
                    images.append({'page':n,'data_url':'data:image/jpeg;base64,'+base64.b64encode(buf.getvalue()).decode()})
    elif suffix=='.pptx':
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist())>5000 or sum(i.file_size for i in archive.infolist())>100*1024*1024:raise ValueError('Reference archive too large')
        deck=Presentation(io.BytesIO(data))
        if not 1<=len(deck.slides)<=80:raise ValueError('Reference must have 1 to 80 slides')
        pages=[{'page':i+1,'text':'\n'.join(s.text for s in sl.shapes if s.has_text_frame)} for i,sl in enumerate(deck.slides)]
    else:raise ValueError('Only PDF/PPTX references supported')
    if sum(len(p['text']) for p in pages)>400000:raise ValueError('Reference text too large')
    return {'pages':pages,'images':images}
