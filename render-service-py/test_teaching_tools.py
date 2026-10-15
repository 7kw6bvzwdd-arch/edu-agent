import io
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from app import app
from test_render import assert_inside_slide

@pytest.mark.parametrize('diagram,expected',[
 ({'kind':'triangle','title':'3、4、5三角形','sides':[3,4,5]},'5'),
 ({'kind':'forces','title':'木块受力','object':'静止木块','unit':'N','equilibrium':True,'forces':[{'label':'重力','x':0,'y':-10},{'label':'支持力','x':0,'y':10}]},'合力'),
 ({'kind':'circuit','title':'并联电路','voltage':6,'resistances':[10,20],'connection':'parallel','closed':True},'0.9'),
])
def test_scientific_diagrams_export_actual_editable_shapes(diagram,expected):
    response=TestClient(app).post('/render/pptx',json={'title':'专业图示','theme':'ocean','slides':[{'title':diagram['title'],'bullets':['观察图示，说明数量关系。'],'scientific':diagram}]})
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    shapes=ppt.slides[0].shapes
    assert any(s.name.startswith('scientific-') for s in shapes),'diagram silently omitted'
    assert expected in '\n'.join(s.text for s in shapes if s.has_text_frame)
    assert not any(s.shape_type==13 for s in shapes)
    if diagram['kind']=='circuit' and diagram['connection']=='parallel':
        assert not any(s.name=='scientific-wire-bottom' for s in shapes),'bare bottom wire would short the parallel resistors'

def test_bad_geometry_and_false_equilibrium_are_rejected():
    for diagram in [{'kind':'triangle','title':'不成立','sides':[1,2,3]}, {'kind':'forces','title':'错误平衡','object':'木块','unit':'N','equilibrium':True,'forces':[{'label':'重力','x':0,'y':-10}]}]:
        result=TestClient(app).post('/render/pptx',json={'title':'课件','slides':[{'title':'图示','scientific':diagram}]})
        assert result.status_code==422

def test_collinear_three_force_diagram_is_readable_and_proportional():
    # Actual buoyancy output: two upward forces shared a tiny shaft and labels.
    diagram={'kind':'forces','title':'三力平衡','object':'金属块','unit':'N','equilibrium':True,'forces':[{'label':'重力G','x':0,'y':-27},{'label':'浮力F浮','x':0,'y':10},{'label':'绳拉力F拉','x':0,'y':17}]}
    response=TestClient(app).post('/render/pptx',json={'title':'浮力','slides':[{'title':'三力平衡','layout':'concept','bullets':['静止金属块受哪些力？'],'scientific':diagram}]})
    assert response.status_code==200,response.text
    slide=Presentation(io.BytesIO(response.content)).slides[0]
    arrows=[next(s for s in slide.shapes if s.name==f'scientific-force-{i}') for i in range(3)]
    assert max(s.top+s.height for s in arrows)-min(s.top for s in arrows)>914400*1.2
    assert arrows[0].height/arrows[1].height==pytest.approx(27/10,rel=1e-5)
    labels=[s for s in slide.shapes if s.has_text_frame and any(s.text.startswith(f) for f in ['重力G','浮力F浮','绳拉力F拉'])]
    assert len(labels)==3
    for i,a in enumerate(labels):
        assert a.text_frame.paragraphs[0].font.size.pt>=14
        for b in labels[i+1:]:
            assert a.left+a.width<=b.left or b.left+b.width<=a.left or a.top+a.height<=b.top or b.top+b.height<=a.top

def test_quadratic_graph_and_inverse_domain():
    graph={'title':'y=x²','x_min':-3,'x_max':3,'y_min':-2,'y_max':10,'series':[{'kind':'quadratic','a':1,'b':0,'c':0,'mark_x':[2]}]}
    result=TestClient(app).post('/render/pptx',json={'title':'函数','slides':[{'title':'抛物线','graph':graph}]})
    assert result.status_code==200,result.text
    shapes=Presentation(io.BytesIO(result.content)).slides[0].shapes
    assert any(s.name.startswith('function-curve-') for s in shapes)
    assert '(2, 4)' in '\n'.join(s.text for s in shapes if s.has_text_frame)
    graph['series']=[{'kind':'inverse','k':2,'mark_x':[0]}]
    assert TestClient(app).post('/render/pptx',json={'title':'函数','slides':[{'title':'反比例','graph':graph}]}).status_code==422

def test_classroom_answer_stays_in_notes_and_formula_is_rendered():
    result=TestClient(app).post('/render/pptx',json={'title':'练习','slides':[{'title':'先独立计算','layout':'exercise','bullets':['计算2+3'],'answer':'答案是5'}]})
    slide=Presentation(io.BytesIO(result.content)).slides[0]
    assert '答案是5' not in '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
    assert '答案是5' in slide.notes_slide.notes_text_frame.text
    response=TestClient(app).post('/render/formula',json={'formula':r'F=\rho g V'})
    assert response.status_code==200 and response.content.startswith(b'\x89PNG')

def test_reference_pdf_contains_real_sampled_images():
    from pypdf import PdfWriter
    writer=PdfWriter()
    for _ in range(8):writer.add_blank_page(width=720,height=405)
    buf=io.BytesIO();writer.write(buf)
    response=TestClient(app).post('/reference/extract',files={'file':('reference.pdf',buf.getvalue(),'application/pdf')})
    assert response.status_code==200,response.text
    data=response.json()
    assert len(data['pages'])==8 and len(data['images'])==6
    assert data['images'][0]['page']==1 and data['images'][-1]['page']==8
    from PIL import Image
    import base64
    image=Image.open(io.BytesIO(base64.b64decode(data['images'][0]['data_url'].split(',')[1])))
    assert image.width==1280

def test_actual_courseware_math_and_flow_fit_office_text_bounds(tmp_path):
    import json, os, subprocess
    from pathlib import Path
    if os.name!='nt' or os.environ.get('TEST_OFFICE_QUALITY')!='1':
        pytest.skip('explicit installed PowerPoint opt-in required')
    source=Path(__file__).parent/'fixtures/courseware-office-regression.json'
    response=TestClient(app).post('/render/pptx',json=json.loads(source.read_text(encoding='utf-8')))
    assert response.status_code==200,response.text
    pptx=tmp_path/'regression.pptx';report=tmp_path/'bounds.json'
    pptx.write_bytes(response.content)
    result=subprocess.run(['powershell','-ExecutionPolicy','Bypass','-File',str(Path(__file__).parents[1]/'scripts/check-courseware-text-bounds.ps1'),'-Path',str(pptx),'-Report',str(report)],capture_output=True,timeout=60)
    assert result.returncode==0,report.read_text(encoding='utf-8-sig') if report.exists() else result.stderr.decode(errors='replace')
    assert json.loads(report.read_text(encoding='utf-8-sig'))=={'slides':3,'overflows':[]}


@pytest.mark.parametrize('case', ['vertex-sidebar', 'density-fraction'])
def test_actual_native_equations_fit_office(tmp_path, case):
    import json, os, subprocess
    from pathlib import Path
    if os.name!='nt' or os.environ.get('TEST_OFFICE_QUALITY')!='1':
        pytest.skip('explicit installed PowerPoint opt-in required')
    data={'title':'二次函数','theme':'ocean','editable_formulas':True,'slides':[{
        'title':'从顶点式读出顶点与对称轴','layout':'concept',
        'bullets':['上节配方结果：y=x²-2x-3=(x-1)²-4','由(x-1)²读对称轴x=1，由-4读最低点纵坐标'],
        'focus':'顶点式中的1和-4在图像上分别对应什么？',
        'formula':r'y=(x-1)^{2}-4,\quad x_{\text{顶点}}=-\frac{b}{2a}=1',
        'graph':{'title':'最低点与对称轴','x_min':-3,'x_max':5,'y_min':-6,'y_max':6,
                 'series':[{'kind':'quadratic','a':1,'b':-2,'c':-3,'mark_x':[1]}]}}]}
    if case=='density-fraction':
        data['title']='浮力'
        data['slides']=[{'title':'分层练习：反求液体密度','layout':'exercise',
            'bullets':['同一物块：空气中示数5 N，全浸没某液体中示数4.2 N（按题设计算）',
                       '物块体积V=2×10⁻⁴ m³，g取10 N/kg，求该液体密度',
                       '写出三步：求F浮→确认V排→代入F浮=ρ液gV排'],
            'focus':'只给两次示数，能求出液体密度吗？需要哪些中间量？',
            'formula':r'\rho_{\text{液}}=\frac{F_{\text{浮}}}{gV_{\text{排}}}'}]
    response=TestClient(app).post('/render/pptx',json=data)
    assert response.status_code==200,response.text
    path=tmp_path/'vertex.pptx';report=tmp_path/'bounds.json';path.write_bytes(response.content)
    result=subprocess.run(['powershell','-ExecutionPolicy','Bypass','-File',str(Path(__file__).parents[1]/'scripts/check-courseware-text-bounds.ps1'),'-Path',str(path),'-Report',str(report)],capture_output=True,timeout=60)
    assert result.returncode==0,report.read_text(encoding='utf-8-sig') if report.exists() else result.stderr.decode(errors='replace')
    assert json.loads(report.read_text(encoding='utf-8-sig'))=={'slides':1,'overflows':[]}
