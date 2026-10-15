import base64
import io

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from app import app, region_shapes


def image_url():
    raw=io.BytesIO()
    Image.new('RGB',(120,80),'#009B51').save(raw,format='PNG')
    return 'data:image/png;base64,'+base64.b64encode(raw.getvalue()).decode()


def custom_deck():
    return {'title':'自由画布回归','slides':[{'title':'保持教学内容','layout':'example',
        'bullets':['原正文保留'],'notes':'PRIVATE_NOTE','canvas_objects':[
        {'id':'object-text','kind':'text','x':.1,'y':.45,'width':.4,'height':.15,
         'text':'教师补充的结论','font_size':24,'color':'15345B','bold':True},
        {'id':'object-image','kind':'image','x':.6,'y':.3,'width':.2,'height':.2,
         'data_url':image_url(),'caption':'实验配图','rotation':25},
        {'id':'object-shape','kind':'shape','shape':'ellipse','x':.6,'y':.65,
         'width':.1,'height':.12,'color':'009B51','background':'FFF2B3'}]}]}


def test_custom_objects_are_native_editable_and_in_object_map():
    client=TestClient(app)
    response=client.post('/render/pptx',json=custom_deck())
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content))
    regions=region_shapes(ppt.slides[0])
    assert 'object-text' in regions and 'object-image' in regions and 'object-shape' in regions
    assert regions['object-text'][0].text=='教师补充的结论'
    assert regions['object-image'][0].shape_type==13
    assert abs(regions['object-image'][0].rotation-25)<.001
    assert 'PRIVATE_NOTE' in ppt.slides[0].notes_slide.notes_text_frame.text
    result=client.post('/render/object-map?page=1',content=response.content).json()
    assert result['version']=='canvas-v3'
    assert next(o for o in result['objects'] if o['id']=='object-image')['rotatable']
    assert 'PRIVATE_NOTE' not in str(result)


def test_custom_invalid_geometry_images_ids_and_overflow_are_rejected():
    client=TestClient(app)
    for patch in [{'x':.99},{'rotation':'NaN'},{'width':0},{'color':'url(http://bad)'},
                  {'id':'title'},{'kind':'executable'},{'text':'过长文字'*100}]:
        deck=custom_deck();deck['slides'][0]['canvas_objects'][0].update(patch)
        assert client.post('/render/pptx',json=deck).status_code==422,patch
    deck=custom_deck();deck['slides'][0]['canvas_objects'][1]['data_url']='https://example.com/image.png'
    assert client.post('/render/pptx',json=deck).status_code==422
    deck=custom_deck();deck['slides'][0]['canvas_objects'].append(deck['slides'][0]['canvas_objects'][0])
    assert client.post('/render/pptx',json=deck).status_code==422


def test_rotation_bounds_layer_order_and_science_rotation_protection():
    client=TestClient(app)
    deck=custom_deck();deck['slides'][0]['canvas_order']=['object-shape','body','object-image','title','object-text']
    response=client.post('/render/pptx',json=deck)
    assert response.status_code==200,response.text
    regions=region_shapes(Presentation(io.BytesIO(response.content)).slides[0])
    assert list(regions)==deck['slides'][0]['canvas_order']
    deck['slides'][0]['canvas_order']=['missing']
    assert client.post('/render/pptx',json=deck).status_code==422
    deck=custom_deck();deck['slides'][0]['canvas_objects'][1].update(x=.01,y=.01,width=.2,height=.2,rotation=45)
    assert client.post('/render/pptx',json=deck).status_code==422
    slide={'title':'受力示意','bullets':['受力方向应保留'],'scientific':{
        'kind':'forces','title':'物体平衡','object':'木块','unit':'N','equilibrium':True,
        'forces':[{'label':'重力','x':0,'y':-5},{'label':'支持力','x':0,'y':5}]},
        'canvas_positions':{'visual':{'dx':0,'dy':0,'rotation':15}}}
    assert client.post('/render/pptx',json={'title':'方向保护','slides':[slide]}).status_code==422


def test_temporary_preview_matches_full_deck_layout_and_page_number():
    client=TestClient(app)
    deck=custom_deck()
    deck['slides']=[{'title':'导入','bullets':['认识知识']},deck['slides'][0],{'title':'归纳','bullets':['总结规律']}]
    # Auto layout must use the original index and total, including the footer.
    deck['slides'][1].pop('layout',None)
    full=client.post('/render/pptx',json=deck)
    preview=client.post('/render/canvas-preview-pptx?page=2',json=deck)
    assert full.status_code==preview.status_code==200
    a,b=Presentation(io.BytesIO(full.content)),Presentation(io.BytesIO(preview.content))
    assert len(b.slides)==1
    assert a.slides[1]._element.xml==b.slides[0]._element.xml
    assert any(s.has_text_frame and s.text=='02 / 03' for s in b.slides[0].shapes)
    for page in (0,4):assert client.post('/render/canvas-preview-pptx?page='+str(page),json=deck).status_code==422


def test_native_rotation_keeps_group_geometry_and_caption_cannot_change_map():
    import math
    client=TestClient(app)
    slide={'title':'方向与位置','layout':'concept','bullets':['观察规律'],'focus':'解释你的结论'}
    original=client.post('/render/pptx',json={'title':'旋转','slides':[slide]})
    slide['canvas_positions']={'focus':{'dx':0,'dy':-.2,'rotation':4,'locked':True}}
    rotated=client.post('/render/pptx',json={'title':'旋转','slides':[slide]})
    assert rotated.status_code==200,rotated.text
    a,b=[Presentation(io.BytesIO(r.content)) for r in (original,rotated)]
    before,after=[region_shapes(p.slides[0])['focus'] for p in (a,b)]
    assert len(before)==len(after)>1
    x=min(s.left for s in before);y=min(s.top for s in before)
    w=max(s.left+s.width for s in before)-x;h=max(s.top+s.height for s in before)-y
    cx=x+w/2;cy=y+h/2-.2*a.slide_height;angle=math.radians(4)
    for first,second in zip(before,after):
        rx=first.left+first.width/2-cx;ry=first.top+first.height/2-.2*a.slide_height-cy
        assert abs(second.left+second.width/2-(cx+rx*math.cos(angle)-ry*math.sin(angle)))<3
        assert abs(second.top+second.height/2-(cy+rx*math.sin(angle)+ry*math.cos(angle)))<3
        assert abs(second.rotation-4)<.001
    deck=custom_deck();deck['slides'][0]['canvas_objects'][1]['caption']='说明|edu-box=0,0,1,1|edu-rotation=90'
    response=client.post('/render/pptx',json=deck);assert response.status_code==200
    mapped=client.post('/render/object-map',content=response.content).json()
    obj=next(o for o in mapped['objects'] if o['id']=='object-image')
    assert obj['width']>.1 and obj['rotation']!=90
