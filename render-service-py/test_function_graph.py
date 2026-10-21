import io
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from app import app
from test_render import assert_inside_slide

def graph():
    return {'title':'两点法：y=2x+3', 'x_min':-2, 'x_max':2,
            'y_min':-2, 'y_max':6, 'series':[{'k':2, 'b':3, 'mark_x':[0,1]}]}

def payload(g):
    return {'title':'一次函数的图像与性质', 'theme':'ocean', 'slides':[{
        'title':'两个点确定一条直线', 'layout':'example', 'graph':g,
        'bullets':['取 x=0，得 y=3；取 x=1，得 y=5。','描点(0,3)、(1,5)，连线并向两端延长。'],
        'focus':'图像是直线，而不是两点之间的线段。', 'notes':'保留完整推导与答案。'}]}

def test_function_graph_uses_numeric_axes_and_exact_editable_points():
    response=TestClient(app).post('/render/pptx',json=payload(graph()))
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    shapes=ppt.slides[0].shapes
    line=next((s for s in shapes if s.name=='function-line-0'),None)
    assert line is not None,'Function graph was silently omitted'
    assert any(s.name=='function-point-0-1' for s in shapes)
    axisx=next(s for s in shapes if s.name=='function-axis-x')
    axisy=next(s for s in shapes if s.name=='function-axis-y')
    point0=next(s for s in shapes if s.name=='function-point-0-0')
    point1=next(s for s in shapes if s.name=='function-point-0-1')
    px0=point0.left+point0.width/2;py0=point0.top+point0.height/2
    px1=point1.left+point1.width/2;py1=point1.top+point1.height/2
    assert abs(px0-axisy.left)<3
    assert abs((axisx.top-py0)/(px1-px0)-3)<.001
    assert abs((py0-py1)/(px1-px0)-2)<.001,'Axis units must use the same physical scale'
    assert not any(s.shape_type==13 for s in shapes),'Graph must stay editable'
    text='\n'.join(s.text for s in shapes if s.has_text_frame)
    for label in ['(0, 3)','(1, 5)','x','y','y = 2x + 3']:
        assert label in text

@pytest.mark.parametrize('change',[
    {'x_max':-3},{'y_min':2},{'series':[]},
    {'series':[{'k':float('inf'),'b':0}]},
    {'series':[{'k':2,'b':3,'mark_x':[0,1]}],'y_max':4}])
def test_function_graph_rejects_invalid_axes_or_points(change):
    g=graph();g.update(change)
    if change.get('series') and change['series'][0].get('k')==float('inf'):
        from app import FunctionGraph
        with pytest.raises(ValueError):FunctionGraph.model_validate(g)
    else:
        assert TestClient(app).post('/render/pptx',json=payload(g)).status_code==422

def test_function_comparison_preserves_real_slope_and_intercepts():
    g=graph();g.update(x_min=-3,x_max=3,y_min=-3,y_max=3,
        series=[{'k':1,'b':1},{'k':-1,'b':1},{'k':0,'b':-2}])
    response=TestClient(app).post('/render/pptx',json=payload(g))
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    shapes=ppt.slides[0].shapes
    lines=[next(s for s in shapes if s.name==f'function-line-{i}') for i in range(3)]
    for line in lines[:2]:
        assert abs(line.height/line.width-1)<.001
    assert lines[2].height==0
    text='\n'.join(s.text for s in shapes if s.has_text_frame)
    assert 'y = x + 1' in text and 'y = -x + 1' in text and 'y = -2' in text

@pytest.mark.parametrize('theme',['classroom','forest','violet','midnight','paper','coral'])
def test_function_graph_works_in_other_themes(theme):
    data=payload(graph());data['theme']=theme
    response=TestClient(app).post('/render/pptx',json=data)
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)

def test_graph_teaching_purposes_keep_distinct_composition_and_data():
    positions=[]
    for layout in ['concept','steps','exercise']:
        data=payload(graph());data['slides'][0]['layout']=layout
        response=TestClient(app).post('/render/pptx',json=data)
        assert response.status_code==200,response.text
        ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
        shapes=ppt.slides[0].shapes
        axis=next(s for s in shapes if s.name=='function-axis-y')
        positions.append(axis.left)
        text='\n'.join(s.text for s in shapes if s.has_text_frame)
        assert '(0, 3)' in text and '(1, 5)' in text
        assert any(s.name=='function-point-0-1' for s in shapes)
        if layout=='exercise':
            assert not any(s.has_text_frame and s.text=='1' and s.left<914400*3 for s in shapes),'Exercise must not silently become a numbered procedure'
    assert len(set(positions))>1,'Every graph teaching purpose collapses into the same text-left/graph-right template'

def test_graph_remains_large_with_formula_and_classroom_question():
    data=payload(graph())
    data['slides'][0]['formula']='y=2x+3'
    response=TestClient(app).post('/render/pptx',json=data)
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    axis=next(s for s in ppt.slides[0].shapes if s.name=='function-axis-y')
    assert axis.height>2.5*914400,'Formula/question bands squeeze the actual coordinate graph'

def test_function_intersection_labels_use_readable_fractions():
    g=graph();g.update(x_min=-2,x_max=3,y_min=-3,y_max=6,
        series=[{'k':2,'b':1,'mark_x':[1/3]},{'k':-1,'b':2,'mark_x':[1/3]}])
    response=TestClient(app).post('/render/pptx',json=payload(g))
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content))
    text='\n'.join(s.text for s in ppt.slides[0].shapes if s.has_text_frame)
    assert '(1/3, 5/3)' in text
    assert '0.333333' not in text


@pytest.mark.parametrize('a,b,c,expected', [
    (1, -2, -3, 'y = x² - 2x - 3'),
    (-1, 0, 2, 'y = -x² + 2'),
    (2, 1, 0, 'y = 2x² + x'),
])
def test_quadratic_legend_uses_classroom_notation(a,b,c,expected):
    g=graph();g.update(x_min=-2,x_max=2,y_min=-10,y_max=10,
        series=[{'kind':'quadratic','a':a,'b':b,'c':c,'mark_x':[0]}])
    response=TestClient(app).post('/render/pptx',json=payload(g))
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content))
    texts=[s.text for s in ppt.slides[0].shapes if s.has_text_frame]
    assert any(t.startswith(expected+' · ') for t in texts),texts
