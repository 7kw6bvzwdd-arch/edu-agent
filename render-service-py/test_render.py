import io, zipfile
import warnings
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from app import app, Deck, render

THEMES = ['forest', 'violet', 'ocean', 'midnight', 'paper', 'coral', 'classroom']
LAYOUTS = ['cover', 'concept', 'steps', 'comparison', 'example', 'exercise', 'summary']


def test_inline_tex_font_wrapper_does_not_print_as_literal_command():
    source = r'y=x²-2x+\mathbf{1}-1-3=(x-1)²-4'
    response = TestClient(app).post('/render/pptx', json={'title': '配方', 'slides': [
        {'title': '配方', 'layout': 'steps', 'bullets': [source], 'notes': '教师原文保留：'+source}]})
    assert response.status_code == 200, response.text
    ppt = Presentation(io.BytesIO(response.content))
    visible = '\n'.join(s.text for s in ppt.slides[0].shapes if s.has_text_frame)
    assert r'\mathbf' not in visible
    assert 'y=x²-2x+1-1-3=(x-1)²-4' in visible
    assert source in ppt.slides[0].notes_slide.notes_text_frame.text


@pytest.mark.parametrize('command,replacement', [('implies','Rightarrow'),('iff','Leftrightarrow'),('impliedby','Leftarrow')])
def test_formula_logical_aliases_preserve_meaning(command, replacement):
    def picture(expr):
        r=TestClient(app).post('/render/pptx',json={'title':'勾股定理','slides':[
            {'title':'面积法证明','layout':'steps','bullets':['用两种方法计算面积'],
             'formula':r'(a+b)^2=4\times(\frac{1}{2}ab)+c^2 '+expr+r' a^2+b^2=c^2'}]})
        assert r.status_code==200,r.text
        ppt=Presentation(io.BytesIO(r.content));assert_inside_slide(ppt)
        return next(s.image.blob for s in ppt.slides[0].shapes if s.shape_type==13)
    assert picture('\\'+command)==picture('\\'+replacement)


@pytest.mark.parametrize('theme', THEMES)
def test_three_step_proof_with_diagram_and_formula_fits(theme):
    # Minimized from the failed 12-page lesson: a narrow body plus formula band.
    slide={'title':'严谨定理证明','layout':'steps','diagram':'triangle',
        'bullets':['用4个全等直角三角形拼成中空的大正方形',
                   '大正方形面积=内部小正方形面积+4个直角三角形面积',
                   '代数展开化简，最终推导出a²+b²=c²完成证明'],
        'formula':r'(a+b)^2=c^2+4 \times \frac{1}{2}ab \Rightarrow a^2+b^2=c^2',
        'notes':'原始讲解备注必须完整保留'}
    r=TestClient(app).post('/render/pptx',json={'title':'勾股定理','theme':theme,'slides':[slide]})
    assert r.status_code==200,r.text
    ppt=Presentation(io.BytesIO(r.content));assert_inside_slide(ppt)
    texts=[s for s in ppt.slides[0].shapes if s.has_text_frame]
    for item in slide['bullets']:
        textbox=next(s for s in texts if s.text==item)
        assert all((run.font.size or p.font.size).pt>=16 for p in textbox.text_frame.paragraphs for run in p.runs)
    assert ppt.slides[0].notes_slide.notes_text_frame.text==slide['notes']


@pytest.mark.parametrize('bad,code,field', [
    ({'layout':'steps','diagram':'triangle','bullets':['密集内容'*35]*6},'text_overflow','bullets'),
    ({'formula':r'x \unsupportedSecretCommand y'},'formula_invalid','formula'),
    ({'chart':{'labels':['甲','乙'],'values':[1]}},'chart_invalid','chart'),
    ({'formula':'x'*251},'content_invalid','formula'),
])
def test_render_error_identifies_page_without_echoing_input(bad,code,field):
    r=TestClient(app).post('/render/pptx',json={'title':'检查','slides':[
        {'title':'正常页','bullets':['要点']},{'title':'第二页私有标题',**bad}]})
    assert r.status_code==422
    d=r.json()['detail']
    assert isinstance(d,dict),d
    assert (d['page'],d['code'],d['field'])==(2,code,field)
    assert '私有标题' not in r.text and 'unsupportedSecretCommand' not in r.text


@pytest.mark.parametrize('missing_declared_font', [False, True])
def test_live_double_escaped_formula_with_chinese_subscripts(caplog, monkeypatch, missing_declared_font):
    if missing_declared_font:
        # PPTX font declarations need not exist on the raster rendering host.
        monkeypatch.setattr('app.FONT', 'Definitely Absent CJK Font')
        monkeypatch.setattr('app._MATHTEXT_FONT', None)
    from PIL import Image
    images = []
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        for formula in [r'F_{浮}=\\rho_{液}gV_{排}', r'F_{浮}=\rho_{液}gV_{排}']:
            deck = Deck(title='浮力', slides=[{'title': '阿基米德原理', 'formula': formula}])
            ppt = Presentation(io.BytesIO(render(deck)))
            assert_inside_slide(ppt)
            picture = next(s for s in ppt.slides[0].shapes if s.shape_type == 13)
            with Image.open(io.BytesIO(picture.image.blob)) as image:
                assert image.getbbox() is not None
            images.append(picture.image.blob)
            assert deck.slides[0].formula == formula
    assert images[0] == images[1], 'Only duplicated command escaping should change'
    assert not any('glyph' in str(w.message).lower() for w in caught)
    assert not any('glyph' in r.message.lower() for r in caplog.records)


def assert_inside_slide(ppt):
    for slide in ppt.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, shape.name
            assert (shape.width > 0 and shape.height > 0) or (shape.shape_type == 9 and shape.width + shape.height > 0), shape.name
            assert shape.left + shape.width <= ppt.slide_width + 2, shape.name
            assert shape.top + shape.height <= ppt.slide_height + 2, shape.name
        # Text containers must not overlap one another (panels may sit behind them).
        texts = [s for s in slide.shapes if s.has_text_frame and s.text.strip()]
        for i, a in enumerate(texts):
            for b in texts[i + 1:]:
                assert (a.left + a.width <= b.left + 2 or b.left + b.width <= a.left + 2
                        or a.top + a.height <= b.top + 2 or b.top + b.height <= a.top + 2), (a.text, b.text)


@pytest.mark.parametrize('theme', THEMES)
def test_all_layouts_preserve_content_and_bounded_geometry(theme):
    slides = [{'title': f'{layout} · 浮力探究', 'layout': layout, 'section': '第一章 · 力与运动',
               'bullets': ['观察实验中的现象', '提出猜想并说明依据', '记录结果，检查单位'],
               'notes': '[S1] 教材第3页'} for layout in LAYOUTS]
    r = TestClient(app).post('/render/pptx', json={'title': '课堂探究', 'theme': theme, 'slides': slides})
    assert r.status_code == 200, r.text
    ppt = Presentation(io.BytesIO(r.content))
    assert_inside_slide(ppt)
    body_geometries = []
    for slide, source in zip(ppt.slides, slides):
        text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
        for value in [source['title'], source['section'], *source['bullets']]:
            assert value in text
        assert source['notes'] == slide.notes_slide.notes_text_frame.text
        body_geometries.append(tuple((s.left, s.top, s.width, s.height) for s in slide.shapes
                                     if s.has_text_frame and any(b in s.text for b in source['bullets'])))
    assert len(set(body_geometries)) == len(LAYOUTS), 'Layouts must differ beyond labels and color'


def test_themes_have_distinct_geometry_and_twenty_slides_render():
    signatures = []
    for theme in THEMES:
        r = TestClient(app).post('/render/pptx', json={'title': '完整教学课件', 'theme': theme,
            'slides': [{'title': '观察与解释', 'bullets': ['观察现象', '解释原因']} for _ in range(20)]})
        assert r.status_code == 200, r.text
        ppt = Presentation(io.BytesIO(r.content))
        assert len(ppt.slides) == 20
        assert_inside_slide(ppt)
        signatures.append(tuple((s.shape_type, s.left, s.top, s.width, s.height) for s in ppt.slides[0].shapes))
        assert any('20 / 20' in s.text for s in ppt.slides[-1].shapes if s.has_text_frame)
    assert len(set(signatures)) == len(THEMES), 'Themes must change composition, not only colors'


@pytest.mark.parametrize('layout', ['concept', 'steps', 'example', 'summary'])
def test_teaching_layouts_use_open_editable_typography_not_card_grid(layout):
    from app import body, Slide
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    p = Presentation()
    p.slide_width = Inches(13.333)
    s = p.slides.add_slide(p.slide_layouts[6])
    items = ['先观察现象并提出一个可检验的问题', '说明实验条件与依据', '用数据解释结果']
    body(s, Slide(title='教学内容', bullets=items), layout, 'forest', .7, 2.55, 11.9, 4.12)
    assert not [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE], layout
    texts = [sh.text for sh in s.shapes if sh.has_text_frame]
    for item in items:
        assert item in texts, (layout, texts)
    assert_inside_slide(p)


@pytest.mark.parametrize('theme', THEMES)
@pytest.mark.parametrize('layout', LAYOUTS)
def test_layouts_with_editable_visuals_and_formula(theme, layout):
    r = TestClient(app).post('/render/pptx', json={'title': '图形与公式', 'theme': theme, 'slides': [
        {'title': '验证规律', 'layout': layout, 'bullets': ['观察图形', '比较数据'],
         'chart': {'title': '实验记录', 'labels': ['甲', '乙', '丙'], 'values': [1, 2, 3]},
         'formula': r'a^2+b^2=c^2'},
        {'title': '浮力示意', 'layout': layout, 'bullets': ['浮力方向竖直向上'], 'diagram': 'buoyancy'},
        {'title': '几何关系', 'layout': layout, 'bullets': ['辨认直角边与斜边'], 'diagram': 'triangle'}]})
    assert r.status_code == 200, r.text
    ppt = Presentation(io.BytesIO(r.content))
    assert_inside_slide(ppt)
    chart = next(s.chart for s in ppt.slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values) == [1, 2, 3]
    assert any(s.shape_type == 13 for s in ppt.slides[0].shapes), 'Formula image missing'
    assert any(s.auto_shape_type == 35 for s in ppt.slides[1].shapes if s.shape_type == 1), 'Editable arrow missing'
    assert any(s.auto_shape_type == 8 for s in ppt.slides[2].shapes if s.shape_type == 1), 'Editable triangle missing'


@pytest.mark.parametrize('slide', [
    {'layout': 'invented'}, {'section': '章' * 81},
    {'bullets': ['单' * 161]}, {'bullets': ['\n' * 80]},
    {'formula': 'x+' * 120 + 'x'},
    {'chart': {'labels': ['一', '二'], 'values': [1]}},
])
def test_invalid_or_overflowing_content_is_rejected(slide):
    r = TestClient(app).post('/render/pptx', json={'title': '边界', 'slides': [{'title': '检查', **slide}]})
    assert r.status_code == 422


def test_long_titles_wrap_and_formula_stays_in_its_band():
    r = TestClient(app).post('/render/pptx', json={'title': '课' * 100, 'theme': 'midnight', 'slides': [
        {'title': '中文标题' * 17, 'layout': 'concept', 'section': '章节' * 20,
         'bullets': ['比较条件\n解释实验结果', 'consistent_units_' * 5], 'formula': r'F_b=\rho g V'}]})
    assert r.status_code == 200, r.text
    ppt = Presentation(io.BytesIO(r.content))
    assert_inside_slide(ppt)
    picture = next(s for s in ppt.slides[0].shapes if s.shape_type == 13)
    assert picture.top >= Inches(6.1) and picture.top + picture.height < Inches(6.85)

def test_real_pptx_has_editable_chinese_slides_and_notes():
    client=TestClient(app)
    response=client.post('/render/pptx',json={'title':'浮力探究','theme':'forest','slides':[
      {'title':'浮力探究','bullets':['教学目标：用实验解释浮力'],'notes':'[S1] 教材第3页'},
      {'title':'课堂练习','bullets':['浸没后深度增大，浮力如何变化？'],'notes':'待核验'}]})
    assert response.status_code==200
    assert zipfile.is_zipfile(io.BytesIO(response.content))
    p=Presentation(io.BytesIO(response.content))
    assert len(p.slides)==2
    assert any('课堂练习' in s.text for s in p.slides[1].shapes if s.has_text_frame)
    assert '[S1]' in p.slides[0].notes_slide.notes_text_frame.text

def test_text_extraction_reports_source_locations():
    r=TestClient(app).post('/extract',files={'file':('lesson.md','浮力\n\n实验记录'.encode(),'text/markdown')})
    assert r.status_code==200
    assert r.json()['pages'][0]['text']=='浮力\n\n实验记录'

def test_empty_slides_are_rejected():
    assert TestClient(app).post('/render/pptx',json={'title':'空','slides':[]}).status_code==422

def test_dense_page_is_rejected_instead_of_clipping():
    r=TestClient(app).post('/render/pptx',json={'title':'过密页面','slides':[
        {'title':'内容过多','diagram':'triangle','bullets':['很长的中文段落'*20]*6}]})
    assert r.status_code==422

def test_three_themes_keep_chart_and_shapes_editable():
    for theme in ['forest','violet','ocean']:
        r=TestClient(app).post('/render/pptx',json={'title':'图表验收','theme':theme,'slides':[
            {'title':'实验数据','bullets':['水中浮力与排开体积的关系'],
             'chart':{'labels':['100','200','300'],'values':[1,2,3]}},
            {'title':'几何图形','bullets':['直角三角形'],'diagram':'triangle'}]})
        assert r.status_code==200
        ppt=Presentation(io.BytesIO(r.content))
        assert any(s.has_chart for s in ppt.slides[0].shapes)
        assert any(s.shape_type==1 for s in ppt.slides[1].shapes)


def test_ordered_measurements_use_editable_line_chart():
    from pptx.enum.chart import XL_CHART_TYPE
    r=TestClient(app).post('/render/pptx',json={'title':'实验记录','slides':[{
        'title':'相同时间间隔的温度变化', 'bullets':['比较不同时间的温度'],
        'chart':{'style':'line','title':'温度/℃ · 示例数据，非实测',
                 'labels':['0 min','1 min','2 min'],'values':[20,24,28]}}]})
    assert r.status_code==200,r.text
    ppt=Presentation(io.BytesIO(r.content));assert_inside_slide(ppt)
    chart=next(s.chart for s in ppt.slides[0].shapes if s.has_chart)
    assert chart.chart_type==XL_CHART_TYPE.LINE_MARKERS
    assert list(chart.series[0].values)==[20,24,28]


def test_ocean_rules_do_not_cross_slide_titles():
    p=Presentation(io.BytesIO(render(Deck(title='浮力',theme='ocean',slides=[
        {'title':'浮力：物体为什么会变轻？','layout':'cover','bullets':['观察实验']},
        {'title':'浮力方向与作用对象','layout':'concept','bullets':['浮力竖直向上']}]))))
    for slide in p.slides:
        title=next(s for s in slide.shapes if s.has_text_frame and s.text.startswith('浮力'))
        rules=[s for s in slide.shapes if s.width>Inches(2) and s.height<Inches(.1)]
        assert not [s for s in rules if s.top < title.top+title.height and s.top+s.height > title.top]


@pytest.mark.parametrize('layout',['concept','summary'])
@pytest.mark.parametrize('visual',[False,True])
def test_six_short_points_fit_open_layout(layout,visual):
    p=Presentation(io.BytesIO(render(Deck(title='排版边界',slides=[{
        'title':'六个具体要点','layout':layout,'bullets':['甲','乙','丙','丁','戊','己'],
        'diagram':'triangle' if visual else '', 'formula':'a^2+b^2=c^2'}]))))
    assert_inside_slide(p)
