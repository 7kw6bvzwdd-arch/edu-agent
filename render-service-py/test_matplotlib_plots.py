import io
import zipfile
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from PIL import Image
from app import app
from test_render import assert_inside_slide


def scatter():
    return {'title':'电流与电压', 'style':'errorbar', 'labels':['甲','乙','丙'],
            'values':[0.1,0.2,0.5], 'x_values':[1,2,5], 'y_errors':[0.01,0.02,0.03],
            'x_label':'电压 U / V', 'y_label':'电流 I / A', 'source_label':'示意数据，非实测'}


def deck(chart):
    return {'title':'实验数据读图', 'theme':'violet', 'slides':[{
        'title':'读出数据变化', 'layout':'concept', 'bullets':['横坐标使用真实数值间距'], 'chart':chart}]}


def test_numeric_experiment_plot_is_embedded_in_real_pptx(tmp_path):
    response=TestClient(app).post('/render/pptx',json=deck(scatter()))
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    picture=next(s for s in ppt.slides[0].shapes if s.name=='matplotlib-data-plot')
    import json
    description=picture._element.nvPicPr.cNvPr.get('descr')
    saved=json.JSONDecoder().raw_decode(description[description.index('{'):])[0]
    assert saved['x_values']==[1,2,5] and saved['y_errors']==[0.01,0.02,0.03]
    assert saved['source_label']=='示意数据，非实测'
    with Image.open(io.BytesIO(picture.image.blob)) as image:
        assert image.width>=1000 and image.height>=500
        # Inspect exported pixels: U=1,2,5 must have 1:3 horizontal gaps.
        pixels=image.convert('RGB');columns=[]
        for x in range(pixels.width):
            if sum(pixels.getpixel((x,y))==(118,81,188) for y in range(pixels.height))>3:columns.append(x)
        groups=[]
        for x in columns:
            if not groups or x>groups[-1][-1]+1:groups.append([])
            groups[-1].append(x)
        centers=[(g[0]+g[-1])/2 for g in groups]
        assert len(centers)==3
        assert (centers[2]-centers[1])/(centers[1]-centers[0])==pytest.approx(3,rel=.02)
    (tmp_path/'experiment.pptx').write_bytes(response.content)


@pytest.mark.parametrize('change',[
    {'x_values':[1,2]}, {'y_errors':[-0.1,0.1,0.1]}, {'y_errors':[0.1]},
    {'x_label':''}, {'source_label':''}, {'x_values':[1,float('inf'),5]},
    {'style':'line'}, {'style':'python'},
])
def test_experiment_plot_rejects_bad_values_before_creating_a_pptx(change):
    chart=scatter();chart.update(change)
    # JSON has no Infinity value: a large exponent still exercises finite validation.
    import json
    raw=json.dumps(deck(chart)).replace('Infinity','1e999')
    response=TestClient(app).post('/render/pptx',content=raw,headers={'Content-Type':'application/json'})
    assert response.status_code==422,response.status_code


def test_function_matplotlib_and_native_chart_can_coexist_in_one_deck():
    g={'renderer':'matplotlib','title':'二次函数 y=x²', 'x_min':-3,'x_max':3,
       'y_min':-1,'y_max':9,'series':[{'kind':'quadratic','a':1,'mark_x':[2]}]}
    body=deck({'title':'分类统计','labels':['甲','乙'],'values':[2,3],'style':'bar'})
    body['slides'].append({'title':'抛物线','graph':g})
    response=TestClient(app).post('/render/pptx',json=body)
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    assert any(s.has_chart for s in ppt.slides[0].shapes)
    assert any(s.name=='matplotlib-function-plot' for s in ppt.slides[1].shapes)
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert 'ppt/charts/chart1.xml' in archive.namelist()
