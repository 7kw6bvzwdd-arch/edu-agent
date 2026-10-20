import base64
import io

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from app import app
from test_render import assert_inside_slide


@pytest.mark.parametrize('count', [5, 6])
@pytest.mark.parametrize('with_picture', [False, True])
def test_history_timeline_preserves_all_events(count, with_picture):
    labels = ['鸦片战争1840—1842', '第二次鸦片战争1856—1860',
              '甲午中日战争1894—1895', '戊戌变法1898', '辛亥革命1911', '中华民国成立1912'][:count]
    slide = {'title': '一条时间轴看近代史主线', 'bullets': ['1840—1911：五个主干事件串成一条主线'],
             'figure': {'kind': 'timeline', 'title': '近代史主干事件顺序',
                        'items': [{'label': label, 'detail': '1912年中华民国成立' if i == 4 else ''}
                                  for i, label in enumerate(labels)]}}
    if with_picture:
        buf = io.BytesIO()
        Image.new('RGB', (300, 200), 'gray').save(buf, format='PNG')
        slide['picture'] = {'data_url': 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode(),
                            'caption': '测试占位图片，非史料', 'source': '测试生成'}
    response = TestClient(app).post('/render/pptx', json={'title': '近代史', 'slides': [slide]})
    assert response.status_code == 200, response.text
    ppt = Presentation(io.BytesIO(response.content))
    assert len(ppt.slides) == 1
    assert_inside_slide(ppt)
    text = '\n'.join(shape.text for shape in ppt.slides[0].shapes if shape.has_text_frame).replace('\u00a0', ' ')
    for label in labels:
        assert label in text
    assert '1912年中华民国成立' in text


@pytest.mark.parametrize('kind,count', [('timeline', 7), ('flow', 5), ('route', 5), ('interval', 5)])
def test_other_figure_limits_remain_bounded(kind, count):
    response = TestClient(app).post('/render/pptx', json={'title': '边界检查', 'slides': [
        {'title': '图解', 'figure': {'kind': kind, 'items': [{'label': '节点'}] * count}}]})
    assert response.status_code == 422
