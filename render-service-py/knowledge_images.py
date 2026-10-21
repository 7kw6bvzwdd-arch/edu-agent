"""Extract bounded embedded figures from trusted uploads, without network or code."""
import base64
import io

from PIL import Image, ImageOps


def extract_images(data, suffix, pages):
    images, warnings = [], []
    total = 0

    def add(raw, page, caption):
        nonlocal total
        if len(raw) > 8 * 1024 * 1024:
            warnings.append(f'第{page}页图片超过8MB，已跳过')
            return
        try:
            with Image.open(io.BytesIO(raw)) as source:
                if min(source.size) < 64 or source.width * source.height > 16000000:
                    return
                source.load()
                image = ImageOps.exif_transpose(source)
                image.thumbnail((2048, 2048))
                output = io.BytesIO()
                image.save(output, 'PNG')
                blob = output.getvalue()
            if len(blob) > 8 * 1024 * 1024 or total + len(blob) > 12 * 1024 * 1024:
                warnings.append('图片提取总量上限12MB，部分图片已跳过')
                return
            total += len(blob)
            images.append(dict(page=page, caption=caption[:600],
                               data_url='data:image/png;base64,' + base64.b64encode(blob).decode()))
        except (OSError, ValueError, Image.DecompressionBombError):
            warnings.append(f'第{page}页包含无法读取的图片，已跳过')

    if suffix == '.pdf':
        import fitz
        with fitz.open(stream=data, filetype='pdf') as document:
            if document.is_encrypted or len(document) > 300:
                raise ValueError('仅支持未加密、300页以内PDF')
            for index, page in enumerate(document):
                # ponytail: embedded rasters only; vector grouping needs a separate renderer.
                for info in page.get_images(full=True)[:12]:
                    if len(images) >= 24:
                        warnings.append('每份文档最多提取24张图片，其余保留在原文中')
                        break
                    if info[2] * info[3] > 16000000 or min(info[2], info[3]) < 64:
                        continue
                    try:
                        pixmap = fitz.Pixmap(document, info[0])
                        if info[1]:
                            pixmap = fitz.Pixmap(pixmap, fitz.Pixmap(document, info[1]))
                        if pixmap.colorspace and pixmap.colorspace.n not in (1, 3):
                            pixmap = fitz.Pixmap(fitz.csRGB, pixmap)
                        add(pixmap.tobytes('png'), index + 1, pages[index]['text'])
                    except (ValueError, RuntimeError, fitz.mupdf.FzErrorBase):
                        warnings.append(f'第{index+1}页包含不支持的图片，已跳过')
                if len(images) >= 24:
                    if index + 1 < len(document):warnings.append('每份文档最多提取24张图片，其余保留在原文中')
                    break
    elif suffix == '.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        def pictures(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from pictures(shape.shapes)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    yield shape

        deck = Presentation(io.BytesIO(data))
        for index, slide in enumerate(deck.slides):
            for shape in pictures(slide.shapes):
                if len(images) >= 24:
                    warnings.append('每份文档最多提取24张图片，其余保留在原文中')
                    break
                add(shape.image.blob, index + 1, pages[index]['text'])
            if len(images) >= 24:
                if index + 1 < len(deck.slides):warnings.append('每份文档最多提取24张图片，其余保留在原文中')
                break
    return images, list(dict.fromkeys(warnings))[:12]
