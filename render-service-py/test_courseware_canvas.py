import io

from fastapi.testclient import TestClient
from pptx import Presentation

from app import app


def test_canvas_moves_actual_native_objects_and_preserves_content():
    client = TestClient(app)
    deck = {'title': '画布回归', 'slides': [{'title': '保持标题', 'layout': 'concept',
            'bullets': ['同一页面正文'], 'notes': 'PRIVATE_NOTES', 'answer': 'PRIVATE_ANSWER'}]}
    original = client.post('/render/pptx', json=deck)
    assert original.status_code == 200
    deck['slides'][0]['canvas_positions'] = {'title': {'dx': .02, 'dy': .03}}
    moved = client.post('/render/pptx', json=deck)
    assert moved.status_code == 200, moved.status_code
    before = Presentation(io.BytesIO(original.content))
    after = Presentation(io.BytesIO(moved.content))
    a = next(s for s in before.slides[0].shapes if s.has_text_frame and s.text == '保持标题')
    b = next(s for s in after.slides[0].shapes if s.has_text_frame and s.text == '保持标题')
    assert abs(b.left - a.left - before.slide_width * .02) < 2
    assert abs(b.top - a.top - before.slide_height * .03) < 2
    assert [(s.text, s.width, s.height) for s in before.slides[0].shapes if s.has_text_frame] == [
        (s.text, s.width, s.height) for s in after.slides[0].shapes if s.has_text_frame]
    manifest = client.post('/render/object-map?page=1', content=moved.content)
    assert manifest.status_code == 200, manifest.text
    title = next(o for o in manifest.json()['objects'] if o['id'] == 'title')
    assert abs(title['x'] - b.left / after.slide_width) < 1e-6
    assert 'PRIVATE' not in manifest.text
    assert 'PRIVATE_ANSWER' in after.slides[0].notes_slide.notes_text_frame.text


def test_canvas_rejects_clipping_unknown_regions_and_nonfinite_offsets():
    client = TestClient(app)
    for positions in [{'title': {'dx': .8, 'dy': 0}}, {'unknown': {'dx': 0, 'dy': 0}},
                      {'title': {'dx': 'NaN', 'dy': 0}}, {'visual': {'dx': .1, 'dy': 0}}]:
        result = client.post('/render/pptx', json={'title': '边界', 'slides': [{
            'title': '完整标题', 'bullets': ['保留正文'], 'canvas_positions': positions}]})
        assert result.status_code == 422, result.status_code


def test_canvas_moves_all_native_graph_parts_together():
    client = TestClient(app)
    slide = {'title': '同一函数图', 'layout': 'example', 'bullets': ['观察函数图像'],
             'graph': {'title': '给定直线', 'x_min': -2, 'x_max': 3, 'y_min': -3, 'y_max': 6,
                       'series': [{'kind': 'linear', 'k': 2, 'b': 1, 'mark_x': [0, 1]}]}}
    deck = {'title': '图示移动', 'slides': [slide]}
    original = client.post('/render/pptx', json=deck)
    slide['canvas_positions'] = {'visual': {'dx': -.01, 'dy': -.01}}
    moved = client.post('/render/pptx', json=deck)
    assert moved.status_code == 200, moved.status_code
    before = Presentation(io.BytesIO(original.content))
    after = Presentation(io.BytesIO(moved.content))
    from app import region_shapes
    a = region_shapes(before.slides[0])['visual']
    b = region_shapes(after.slides[0])['visual']
    assert len(a) == len(b) > 10
    for first, second in zip(a, b):
        assert abs(second.left - first.left + before.slide_width * .01) < 2
        assert abs(second.top - first.top + before.slide_height * .01) < 2
        assert (first.width, first.height) == (second.width, second.height)


def test_canvas_scale_keeps_native_graph_geometry_fonts_and_notes():
    from app import region_shapes
    client = TestClient(app)
    slide = {'title': '一次函数图像', 'layout': 'example', 'bullets': ['观察直线的截距和标记点'],
             'notes': 'PRIVATE_TEACHER_NOTE',
             'graph': {'title': 'y = 2x + 1', 'x_min': -2, 'x_max': 3, 'y_min': -3, 'y_max': 6,
                       'series': [{'kind': 'linear', 'k': 2, 'b': 1, 'mark_x': [0, 1]}]}}
    deck = {'title': '等比例缩放回归', 'slides': [slide]}
    original = client.post('/render/pptx', json=deck)
    slide['canvas_positions'] = {'visual': {'dx': -.12, 'dy': -.03, 'scale': 1.1}}
    scaled = client.post('/render/pptx', json=deck)
    assert scaled.status_code == 200, scaled.text
    before, after = [Presentation(io.BytesIO(r.content)) for r in (original, scaled)]
    a, b = [region_shapes(p.slides[0])['visual'] for p in (before, after)]
    x, y = min(i.left for i in a), min(i.top for i in a)
    assert len(a) == len(b) > 10
    for first, second in zip(a, b):
        assert abs(second.left-(x+(first.left-x)*1.1-before.slide_width*.12)) < 3
        assert abs(second.top-(y+(first.top-y)*1.1-before.slide_height*.03)) < 3
        assert abs(second.width-first.width*1.1) < 3
        assert abs(second.height-first.height*1.1) < 3
        fonts1=first._element.xpath('.//a:defRPr/@sz')
        fonts2=second._element.xpath('.//a:defRPr/@sz')
        assert [round(int(v)*1.1) for v in fonts1] == [int(v) for v in fonts2]
    assert 'PRIVATE_TEACHER_NOTE' in after.slides[0].notes_slide.notes_text_frame.text
    manifest=client.post('/render/object-map?page=1', content=scaled.content).json()
    visual=next(o for o in manifest['objects'] if o['id']=='visual')
    assert abs(visual['width']-(max(i.left+i.width for i in b)-min(i.left for i in b))/after.slide_width)<1e-6
    assert 0<visual['min_scale']<=1
    assert manifest['version']=='canvas-v3'


def test_canvas_scale_table_chart_and_editable_math():
    from app import region_shapes
    client=TestClient(app)
    for extra, region in [({'table': {'headers':['时间/s','路程/m'], 'rows':[['1','2'],['2','4']]}},'visual'),
                          ({'chart': {'labels':['A','B'], 'values':[2,4]}},'visual'),
                          ({'scientific':{'kind':'triangle','title':'Side lengths','sides':[3,4,5]}},'visual'),
                          ({'scientific':{'kind':'forces','title':'Balanced forces','object':'Block','unit':'N','equilibrium':True,'forces':[{'label':'Weight','x':0,'y':-5},{'label':'Support','x':0,'y':5}]}},'visual'),
                          ({'scientific':{'kind':'circuit','title':'Series circuit','voltage':6,'resistances':[10,20],'connection':'series','closed':True}},'visual'),
                          ({'formula': r'F=\frac{G}{2}'},'formula')]:
        slide={'title':'原生对象比例调整','layout':'example','bullets':['核对原始数据和对象关系'],**extra}
        deck={'title':'A','editable_formulas':True,'slides':[slide]}
        original=client.post('/render/pptx',json=deck)
        assert original.status_code==200,original.text
        slide['canvas_positions']={region:{'dx':-.03,'dy':-.12 if region=='formula' else -.03,'scale':1.05}}
        scaled=client.post('/render/pptx',json=deck)
        assert scaled.status_code==200,scaled.text
        a,b=[region_shapes(Presentation(io.BytesIO(r.content)).slides[0])[region] for r in (original,scaled)]
        for first,second in zip(a,b):
            assert abs(second.width-first.width*1.05)<3
            assert abs(second.height-first.height*1.05)<3
            if first.has_table:
                assert [round(c.width*1.05) for c in first.table.columns]==[c.width for c in second.table.columns]
                assert [round(row.height*1.05) for row in first.table.rows]==[row.height for row in second.table.rows]
            if first.has_chart:
                assert list(first.chart.series[0].values)==list(second.chart.series[0].values)
                assert abs(second.chart.font.size-first.chart.font.size*1.05)<130
            word='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
            sizes1=[int(n.get(word+'val')) for n in first._element.iter(word+'sz')]
            sizes2=[int(n.get(word+'val')) for n in second._element.iter(word+'sz')]
            assert [round(v*1.05) for v in sizes1]==sizes2


def test_canvas_scale_rejects_invalid_values_clipping_and_unreadable_text():
    client=TestClient(app)
    for scale in [0,-1,.1,1.51,'NaN','Infinity',1.5]:
        result=client.post('/render/pptx',json={'title':'A','slides':[{
            'title':'A','bullets':['A'],'canvas_positions':{'title':{'dx':0,'dy':0,'scale':scale}}}]})
        assert result.status_code==422,(scale,result.status_code)
    result=client.post('/render/pptx',json={'title':'A','slides':[{
        'title':'A','layout':'example','bullets':['A'],'chart':{'labels':['A','B'],'values':[2,4]},
        'canvas_positions':{'visual':{'dx':0,'dy':0,'scale':.8}}}]})
    assert result.status_code==422


def test_canvas_raster_formula_enlarge_can_return_to_original_scale():
    client=TestClient(app)
    deck={'title':'Raster formula','slides':[{'title':'Formula','layout':'concept',
          'bullets':['Explain the relationship'], 'formula':'F=G/2',
          'canvas_positions':{'formula':{'dx':-.03,'dy':-.12,'scale':1.05}}}]}
    rendered=client.post('/render/pptx',json=deck)
    assert rendered.status_code==200,rendered.text
    manifest=client.post('/render/object-map?page=1',content=rendered.content).json()
    obj=next(o for o in manifest['objects'] if o['id']=='formula')
    assert abs(obj['min_scale']*1.05-1)<1e-6


def test_canvas_inline_title_body_and_focus_are_written_to_pptx():
    client=TestClient(app)
    deck={'title':'Inline text','slides':[{'title':'新标题','layout':'example',
        'bullets':['第一条正文','第二条正文'],'focus':'课堂问题：为什么？',
        'notes':'PRIVATE_NOTE'}]}
    response=client.post('/render/pptx',json=deck)
    assert response.status_code==200,response.text
    presentation=Presentation(io.BytesIO(response.content))
    visible='\n'.join(s.text for s in presentation.slides[0].shapes if s.has_text_frame)
    assert '新标题' in visible and '第一条正文' in visible and '第二条正文' in visible
    assert '课堂问题：为什么？' in visible
    assert 'PRIVATE_NOTE' in presentation.slides[0].notes_slide.notes_text_frame.text


def test_canvas_inline_text_rejects_overlong_content_without_rendering_partial_pptx():
    client=TestClient(app)
    for slide in [
        {'title':'x'*71,'bullets':['正文']},
        {'title':'标题','bullets':['正文']*7},
        {'title':'标题','bullets':['x'*161]},
        {'title':'标题','bullets':['正文'],'focus':'x'*61},
    ]:
        response=client.post('/render/pptx',json={'title':'Inline limits','slides':[slide]})
        assert response.status_code==422
