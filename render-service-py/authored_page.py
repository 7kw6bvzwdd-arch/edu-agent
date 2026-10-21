"""Convert isolated browser output to editable PPT text plus artwork."""
import base64
import io
import json
import os
from pathlib import Path
import re
import subprocess
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

def render_authored_page(slide, source, notes, answer='', picture=None):
    bridge = Path(os.environ.get('AUTHORED_PAGE_BRIDGE', Path(__file__).resolve().parent.parent / 'animation-video' / 'authored-page.mjs'))
    try:
        result = subprocess.run([os.environ.get('NODE_BINARY', 'node'), str(bridge)],
            input=json.dumps({'html': source,'picture':picture.data_url if picture else None}), text=True, encoding='utf-8',
            capture_output=True, timeout=18, check=True, creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        data = json.loads(result.stdout)
    except (subprocess.SubprocessError, ValueError, OSError) as exc:
        raise ValueError('authored_page_invalid') from exc
    if data.get('error'):
        allowed = {
            'authored_text_unmarked', 'authored_text_overlap', 'authored_text_outside',
            'authored_text_overflow', 'authored_text_transform', 'authored_text_ids',
            'authored_text_empty', 'authored_text_unmeasurable',
        }
        raise ValueError(data['error'] if data['error'] in allowed else 'authored_page_invalid')
    slide.shapes.add_picture(io.BytesIO(base64.b64decode(data['png'])), 0, 0, width=Inches(13.333333), height=Inches(7.5)).name='authored-artwork'
    for item in data['elements']:
        for i,line in enumerate(item['lines']):
            box=slide.shapes.add_textbox(Inches(line['x']/96), Inches(line['y']/96), Inches((line['width']*1.025+12)/96), Inches((line['height']+4)/96))
            box.name='authored-'+item['id']+('-'+str(i) if i else '')
            box._element.xpath('.//p:cNvPr')[0].set('descr','edu-object=direct:'+item['id']+'|edu-rotatable=1')
            tf=box.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
            p=tf.paragraphs[0];p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(line['height']*.75)
            for value in line['runs']:
                run=p.add_run();run.text=value['text'];font=run.font
                font.name=value['font'].split(',')[0].strip(chr(34)+chr(39)+' ');
                from pptx.oxml.xmlchemy import OxmlElement
                east=OxmlElement('a:ea');east.set('typeface',font.name);run._r.get_or_add_rPr().append(east);font.size=Pt(value['size']*.75);font.bold=value['bold'];font.italic=value['italic']
                colors=re.findall(r'[\d.]+',value['color']);font.color.rgb=RGBColor(*[int(float(v)) for v in colors[:3]])
    slide.notes_slide.notes_text_frame.text=notes+'\n直接生成页面：文字可编辑，复杂图形保留为背景图层。'+ ('\n课堂答案：'+answer if answer else '')+ ('\n配图：'+picture.caption+'；'+picture.source if picture else '')
