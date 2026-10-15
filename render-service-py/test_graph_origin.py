import io
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from app import app


@pytest.mark.parametrize('bounds', [(0, 5, 0, 12), (-5, 0, -12, 0)])
def test_actual_pptx_accepts_origin_on_axis_boundary(bounds):
    xmin,xmax,ymin,ymax=bounds
    deck={'title':'函数范围','theme':'ocean','slides':[{
        'title':'观察第一象限或第三象限','layout':'example',
        'bullets':['按题设计算，非实测；x/y每单位等长。'],
        'graph':{'title':'y=2x','x_min':xmin,'x_max':xmax,'y_min':ymin,'y_max':ymax,
                 'series':[{'kind':'linear','k':2,'b':0,'mark_x':[0]}]}}]}
    with TestClient(app) as client:
        response=client.post('/render/pptx',json=deck)
    assert response.status_code==200,response.text
    presentation=Presentation(io.BytesIO(response.content))
    assert len(presentation.slides)==1
    assert len(presentation.slides[0].shapes)>10
    for shape in presentation.slides[0].shapes:
        assert shape.left>=0 and shape.top>=0
        assert shape.left+shape.width<=presentation.slide_width+1
        assert shape.top+shape.height<=presentation.slide_height+1
