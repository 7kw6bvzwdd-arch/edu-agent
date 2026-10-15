import io

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import app, THEMES, text_lines
from test_render import assert_inside_slide


def test_reading_intersection_exercise_retains_points_without_printing_coordinates():
    data={'title':'读出两条直线的交点','layout':'exercise','bullets':['观察两直线，读出交点坐标'],
          'answer':'第一步：联立两式\n第二步：交点为(1/3,5/3)',
          'graph':{'title':'给定两条直线','x_min':-2,'x_max':3,'y_min':-3,'y_max':6,
                   'series':[{'kind':'linear','k':2,'b':1,'mark_x':[1/3]},
                             {'kind':'linear','k':-1,'b':2,'mark_x':[1/3]}]}}
    response=TestClient(app).post('/render/pptx',json={'title':'一次函数','slides':[data]})
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    visible='\n'.join(s.text for s in ppt.slides[0].shapes if s.has_text_frame)
    assert '1/3' not in visible and '5/3' not in visible
    assert len([s for s in ppt.slides[0].shapes if s.name.startswith('function-point-')])==2
    assert data['answer'] in ppt.slides[0].notes_slide.notes_text_frame.text


def test_scientific_comparison_preserves_two_distinct_comparison_groups():
    data={'title':'漂浮和浸没绳系比较','layout':'comparison','bullets':['漂浮：只有重力和浮力','绳系：重力、浮力和绳拉力'],
          'scientific':{'kind':'forces','title':'漂浮小球受力','object':'小球','unit':'N','equilibrium':True,
                        'forces':[{'label':'重力','x':0,'y':-5},{'label':'浮力','x':0,'y':5}]}}
    response=TestClient(app).post('/render/pptx',json={'title':'浮力','slides':[data]})
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    text='\n'.join(s.text for s in ppt.slides[0].shapes if s.has_text_frame)
    assert '\nA\n' in text and '\nB\n' in text

def test_supplied_graph_coordinates_remain_visible_when_not_the_requested_answer():
    data={'title':'由给定点判断斜率','layout':'exercise','bullets':['图中已标出两个点，请判断函数递增还是递减'],
          'graph':{'title':'给定点及直线','x_min':-2,'x_max':3,'y_min':-3,'y_max':6,
                   'series':[{'kind':'linear','k':2,'b':1,'mark_x':[0,1]}]}}
    response=TestClient(app).post('/render/pptx',json={'title':'一次函数','slides':[data]})
    assert response.status_code==200,response.text
    visible='\n'.join(s.text for s in Presentation(io.BytesIO(response.content)).slides[0].shapes if s.has_text_frame)
    assert '(0, 1)' in visible and '(1, 3)' in visible


def test_chinese_quotation_width_matches_actual_office_wrap():
    # Office uses three lines at 23pt/4.4in for this real misconception prompt.
    assert text_lines('说法A：“b越大，直线越陡”——请用图中两条直线检验',4.4,23)>=3


def test_parallel_lesson_keeps_long_question_and_complete_conditions():
    data={'title':'并联电路：电压相等，干路电流等于支路电流之和','layout':'concept',
          'bullets':['并联各支路两端电压相等，都等于电源电压','电流在节点分支：I=I₁+I₂'],
          'focus':'并联时每个电阻两端电压各是多少？干路电流与支路电流有什么关系？',
          'formula':r'\frac{1}{R}=\frac{1}{R_1}+\frac{1}{R_2},\quad I=I_1+I_2',
          'scientific':{'kind':'circuit','connection':'parallel','closed':True,'resistances':[10,20],
                        'voltage':6,'title':'并联电流与电压规律（理想电源，示意非实测）'}}
    response=TestClient(app).post('/render/pptx',json={'title':'串并联电路','slides':[data]})
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    for text in [*data['bullets'],data['focus']]:
        assert any(text in s.text.replace('\u00a0',' ') for s in ppt.slides[0].shapes if s.has_text_frame)


@pytest.mark.parametrize('layout', ['concept', 'example'])
@pytest.mark.parametrize('connection', ['series', 'parallel'])
@pytest.mark.parametrize('resistances', [[10,20],[1000,200,300]])
def test_circuit_with_question_and_formula_remains_projection_size(layout, connection, resistances):
    data={'title':'电流与电压规律','layout':layout,
          'bullets':['理想电源6 V，电阻10 Ω与20 Ω，分析电流路径',
                     '结合连接方式，说明电流与电压的关系'],
          'focus':'电流在哪处分流？为什么？', 'formula':'I=U/R',
          'notes':'原教师说明：理想计算，非实测。',
          'scientific':{'kind':'circuit','title':'理想电路（示意，非实测）',
                        'voltage':6,'resistances':resistances,
                        'connection':connection,'closed':True}}
    response=TestClient(app).post('/render/pptx',json={'title':'串并联电路','theme':'ocean','slides':[data]})
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    slide=ppt.slides[0]
    wire=next(s for s in slide.shapes if s.name=='scientific-wire-left')
    assert wire.height/914400>=1.5,'footer strips crushed the circuit into a narrow ribbon'
    labels=[s for s in slide.shapes if s.has_text_frame and
            (s.text.startswith(('R1=', 'R2=', 'I1=', 'I2=', '理想电源', '总电阻')))]
    assert len(labels)>=4
    for label in labels:
        assert all(p.font.size.pt>=18 for p in label.text_frame.paragraphs),'essential circuit label too small for projection'
    for text in [data['title'],*data['bullets'],data['focus']]:
        assert any(text in s.text.replace('\u00a0',' ') for s in slide.shapes if s.has_text_frame)
    assert data['notes']==slide.notes_slide.notes_text_frame.text


def test_function_graph_labels_remain_readable_for_classroom_projection():
    data={'title':'两条直线的交点','layout':'concept','bullets':['比较斜率与截距'],
          'graph':{'title':'同一坐标比例比较函数','x_min':-2,'x_max':3,'y_min':-3,'y_max':6,
                   'series':[{'kind':'linear','k':2,'b':1,'mark_x':[1/3]},
                             {'kind':'linear','k':-1,'b':2,'mark_x':[1/3]}]}}
    response=TestClient(app).post('/render/pptx',json={'title':'一次函数','slides':[data]})
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    labels=[s for s in ppt.slides[0].shapes if s.has_text_frame and
            (s.text in ['x','y','0','2','-2'] or s.text.startswith('y = '))]
    assert len(labels)>=4
    for label in labels:
        assert all(p.font.size.pt>=14 for p in label.text_frame.paragraphs),'graph label below the projection reading floor'


def test_actual_buoyancy_force_labels_fit_without_shrinking():
    diagram={'kind':'forces','title':'漂浮小球受力','object':'漂浮在水面的小球','unit':'N','equilibrium':True,
             'forces':[{'label':'重力*G* 5 N','x':0,'y':-5},{'label':'浮力*F*_浮 5 N','x':0,'y':5}]}
    response=TestClient(app).post('/render/pptx',json={'title':'浮力','slides':[{
        'title':'漂浮与绳系浸没','layout':'comparison','bullets':['辨析平衡条件'],
        'scientific':diagram,'focus':'受力平衡能推出全浸没吗？',
        'formula':r'\text{漂浮：}F_{\text{浮}}=G\quad\text{绳系浸没：}F_{\text{浮}}+F_{\text{拉}}=G'}]})
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    for force in diagram['forces']:
        label=next(s for s in ppt.slides[0].shapes if s.has_text_frame and s.text.replace('\u00a0',' ')==force['label'])
        assert label.text_frame.paragraphs[0].font.size.pt>=18


@pytest.mark.parametrize('layout', ['steps','comparison','example','exercise'])
def test_working_pages_do_not_repeat_two_full_width_footer_panels(layout):
    data={'title':'结合证据说明依据','layout':layout,'bullets':['保留已给条件','写出必要推理','说明结论'],
          'focus':'哪一个条件决定结论？','formula':'F=ma'}
    response=TestClient(app).post('/render/pptx',json={'title':'教学任务','theme':'ocean','slides':[data]})
    assert response.status_code==200,response.text
    ppt=Presentation(io.BytesIO(response.content));assert_inside_slide(ppt)
    filled=[]
    for shape in ppt.slides[0].shapes:
        if shape.top/914400>4.5 and shape.width/914400>10 and shape.height/914400>.4:
            if shape.fill.type==1 and str(shape.fill.fore_color.rgb)==THEMES['ocean'][4]:
                filled.append(shape)
    assert not filled,'question and formula repeated as broad color panels instead of task-led typography'


def test_text_only_cover_uses_available_reading_width():
    data={'title':'从一次到二次：怎样发现图象的转折？','layout':'cover',
          'bullets':['先观察同一组函数值，再用图象和配方解释变化'],
          'focus':'函数值在哪个位置开始回升？'}
    response=TestClient(app).post('/render/pptx',json={'title':'二次函数','theme':'ocean','slides':[data]})
    assert response.status_code==200,response.text
    slide=Presentation(io.BytesIO(response.content)).slides[0]
    title=next(s for s in slide.shapes if s.has_text_frame and s.text==data['title'])
    assert title.width/914400>10,'text-only opening kept a narrow column beside an empty visual slot'
