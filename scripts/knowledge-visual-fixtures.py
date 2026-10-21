"""Labelled local fixtures for real PDF/PPTX extraction and click animation."""
import io
import sys
from pathlib import Path

import fitz
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

root = Path(__file__).resolve().parents[1]
destination = root / 'work' / 'knowledge-click-20261020'
destination.mkdir(parents=True, exist_ok=True)
image = Image.new('RGB', (640, 360), '#eef7fc')
draw = ImageDraw.Draw(image)
draw.rectangle((80, 260, 560, 310), fill='#3682a2')
draw.ellipse((200, 40, 440, 120), fill='#acceda')
draw.line([(130, 250), (160, 140), (210, 100)], fill='#277d60', width=8)
draw.line([(420, 110), (460, 240)], fill='#3682a2', width=8)
draw.text((30, 20), 'FIXED TEST FIXTURE - water cycle', fill='#18324c')
raw = io.BytesIO(); image.save(raw, 'PNG')
pdf = fitz.open(); page = pdf.new_page()
page.insert_text((40, 40), 'Water cycle fixed extraction fixture')
page.insert_image(fitz.Rect(40, 70, 550, 357), stream=raw.getvalue())
pdf.save(destination / 'cycle.pdf'); pdf.close()
deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
slide.shapes.add_textbox(Inches(.5), Inches(.2), Inches(7), Inches(.6)).text = 'Water cycle fixed extraction fixture'
slide.shapes.add_picture(io.BytesIO(raw.getvalue()), Inches(1), Inches(1), width=Inches(7))
deck.save(destination / 'cycle.pptx')
sys.path.insert(0, str(root / 'render-service-py'))
from app import Deck, render
click = dict(title='固定样例：原生点击动画', theme='classroom', slides=[dict(
    title='点击依次比较函数图像', bullets=['每次点击出现一条曲线；比较开口方向。'],
    notes='固定工程验收样例，不代表真实AI生成或课堂验收。', graph=dict(
        title='同一坐标系中的二次函数', x_min=-3, x_max=3, y_min=-4, y_max=4,
        reveal_on_click=True, series=[dict(kind='quadratic', a=1, mark_x=[1]),
                                     dict(kind='quadratic', a=-1, mark_x=[-1])]))])
(destination / 'native-click.pptx').write_bytes(render(Deck(**click)))
print(destination)
