"""Generate five real sample decks through the same HTTP flow a teacher uses."""
import io
import json
import time
import urllib.request
from pathlib import Path
from pptx import Presentation
from smoke import request, ROOT, BASE, cfg

assert request('/health')['mode']=='sample'
token=request('/auth/login',{'username':'admin','password':cfg['SEED_PASSWORD']})['token']
results=[]
for index,topic in enumerate(['浮力','勾股定理','一次函数','欧姆定律','光的反射']):
    theme=['forest','violet','ocean'][index%3]
    started=time.monotonic()
    ident=request('/tasks',{'type':'courseware','topic':topic,'subject':'物理' if index in [0,3,4] else '数学','grade':'八年级','textbook':'人教版','pages':8,'theme':theme,'use_kb':True,'idempotency_key':'accept-'+str(time.time_ns())},token)['id']
    for _ in range(100):
        task=request('/tasks/'+ident,token=token)
        if task['status']=='awaiting_confirmation':break
        time.sleep(.2)
    assert task['status']=='awaiting_confirmation',task
    request('/tasks/'+ident+'/confirm',{'result':task['result'],'revision':task['revision'],'accept_review':True},token)
    for _ in range(100):
        task=request('/tasks/'+ident,token=token)
        if task['status'] in ['succeeded','failed']:break
        time.sleep(.2)
    assert task['status']=='succeeded',task
    req=urllib.request.Request(BASE+'/tasks/'+ident+'/download/pptx',headers={'Authorization':'Bearer '+token})
    with urllib.request.urlopen(req) as r:content=r.read()
    deck=Presentation(io.BytesIO(content))
    assert len(deck.slides)==8
    texts=[shape.text for sl in deck.slides for shape in sl.shapes if shape.has_text_frame]
    assert any('例题讲解' in text for text in texts)
    assert any('课堂练习' in text for text in texts)
    assert any(topic in text for text in texts)
    path=ROOT/'work'/f'acceptance-{index+1}.pptx';path.write_bytes(content)
    result={'topic':topic,'theme':theme,'slides':8,'editable_text_shapes':len(texts),'bytes':len(content),'elapsed_seconds':round(time.monotonic()-started,2),'task_id':ident,'file':str(path),'mode':'sample','sources':len(task['sources'])}
    results.append(result)
    print(json.dumps(result),flush=True)
(ROOT/'work/courseware-acceptance.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
